"""Đối soát R1–R4 (lát R1, chỉ đọc; giả định reviewer, chờ đối chiếu BA)."""
from datetime import datetime
from io import BytesIO

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from server.database import Base, get_db
from server.main import app
from server.models import (
    ArrangementDossier,
    ArrangementImport,
    AssignedDocument,
    Project,
    ProjectCase,
    ProjectDocumentAsset,
    ProjectReportUnit,
    Submission,
    Template,
    User,
)
from server.models_scan import CaseScanFile, CaseScanPackage
from server.routers.auth import get_current_user
from server.services.reconciliation_service import build_reconciliation, reconciliation_workbook

ADMIN = {"id": 1, "username": "admin", "full_name": "Quản trị", "role": "admin", "session_id": "s"}
STAFF = {"id": 2, "username": "nv", "full_name": "Nhân viên", "role": "user", "session_id": "s"}


@pytest.fixture()
def db(tmp_path):
    engine = create_engine(f"sqlite:///{(tmp_path / 'reconcile.sqlite3').as_posix()}")
    Base.metadata.create_all(bind=engine)
    session = sessionmaker(bind=engine)()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


@pytest.fixture()
def project(db):
    admin = User(username="admin", password="x", role="admin")
    template = Template(name="T", filename="t.xlsx")
    db.add_all([admin, template])
    db.flush()
    project = Project(
        name="Bộ Y tế", root_folder_name="Goc", template_id=template.id, template_name_snapshot="T",
        template_filename_snapshot="t.xlsx", case_level=1, report_mode="pdf", created_by_user_id=admin.id, status="new",
    )
    db.add(project)
    db.flush()
    cases = {key: ProjectCase(project_id=project.id, case_key=key, display_name=key) for key in ("0001", "0002", "0003", "0004")}
    db.add_all(cases.values())
    db.flush()
    units = {key: ProjectReportUnit(project_id=project.id, case_id=row.id, report_key=key, display_name=key) for key, row in cases.items()}
    db.add_all(units.values())
    dropped_by = ArrangementImport(project_id=project.id, file_name="m2.xlsx", file_sha256="1" * 64)
    db.add(dropped_by)
    db.flush()

    def dossier(box, number, sheets, suffix="", dropped=False):
        db.add(ArrangementDossier(
            project_id=project.id, case_id=cases[box].id, box_number=int(box), dossier_number=number, dossier_suffix=suffix,
            fonds_code="P", fonds_name="Phông", catalog_number="01", title=f"Hồ sơ {number}{suffix}", start_date="01/01/2020",
            end_date="31/12/2020", start_year=2020, maintenance_code="01", sheet_count=sheets, source_row=number,
            missing_from_import_id=dropped_by.id if dropped else None,
        ))

    dossier("0001", 12, 10); dossier("0001", 13, 5, "a"); dossier("0001", 14, 3); dossier("0001", 15, 99, dropped=True)
    dossier("0002", 1, 4)
    dossier("0004", 1, 50)

    def asset(box, path, size, *statuses, asset_status="active", linked=True):
        document = None
        if linked:
            document = AssignedDocument(original_filename=path, uuid_filename=f"u-{path}-{asset_status}")
            db.add(document)
            db.flush()
            for status in statuses:
                db.add(Submission(template_id=template.id, assigned_document_id=document.id, data_json="{}", status=status))
        db.add(ProjectDocumentAsset(
            project_id=project.id, case_id=cases[box].id, report_unit_id=units[box].id,
            assigned_document_id=document.id if document else None, relative_path=path,
            normalized_relative_path=path.casefold(), original_filename=path.rsplit("/", 1)[-1],
            storage_filename=f"s-{path.replace('/', '-')}-{asset_status}", content_sha256="0" * 64, byte_size=size,
            status=asset_status,
        ))

    asset("0001", "0001/012/BIA.pdf", 10, "pending_review")
    asset("0001", "0001/012/1.pdf", 100, "completed")
    asset("0001", "0001/012/2.pdf", 50, "pending_review")
    asset("0001", "0001/012/3.pdf", 70, "completed", asset_status="replaced")
    asset("0001", "0001/013A/1.pdf", 60, "completed", "completed")
    asset("0001", "0001/099/1.pdf", 1)
    asset("0001", "0001/abc/1.pdf", 1, linked=False)
    asset("0001", "0001/loose.pdf", 1, linked=False)
    asset("0002", "0002/001/1.pdf", 200, "completed")
    asset("0003", "0003/001/1.pdf", 300)
    db.flush()

    def package(box, version, status, files):
        row = CaseScanPackage(case_id=cases[box].id, version=version, submitted_by_user_id=admin.id, source_path="x",
                              status=status, finished_at=datetime(2026, 10, 1))
        db.add(row)
        db.flush()
        for path, size, pages in files:
            db.add(CaseScanFile(package_id=row.id, relative_path=path, file_size=size, page_count=pages,
                                is_pdf=path.lower().endswith(".pdf")))

    package("0001", 1, "done", [("012/1.pdf", 1, 1)])
    package("0001", 2, "done", [
        ("012\\BIA.pdf", 10, 1), ("012/1.PDF", 100, 8), ("012/2.pdf", 999, 12), ("013a/1.pdf", 60, 6),
        ("014/1.pdf", 5, 11), ("012/3.pdf", 70, -1), ("ghi_chu.txt", 3, 0),
    ])
    package("0001", 3, "failed", [("012/1.pdf", 100, 500)])
    package("0002", 1, "done", [("001/1.pdf", 200, 20)])
    package("0004", 1, "done", [("001/1.pdf", 1, 10)])
    db.commit()
    return project


def test_r1_catalog_against_dossier_folders(db, project):
    rows = build_reconciliation(db, project_id=project.id)["rows"]["R1"]
    assert rows == [
        [1, "12", "Hồ sơ 12", "012", 3, "Khớp", ""],
        [1, "13a", "Hồ sơ 13a", "013A", 1, "Khớp", ""],
        [1, "14", "Hồ sơ 14", "", 0, "Thiếu thư mục", ""],
        [1, "99", "", "099", 1, "Thừa thư mục", ""],
        [1, "", "", "abc", 1, "Thừa thư mục", "Tên thư mục không đọc được số hồ sơ"],
        [2, "1", "Hồ sơ 1", "001", 1, "Khớp", ""],
        [3, "1", "", "001", 1, "Thừa thư mục", ""],
        [4, "1", "Hồ sơ 1", "", 0, "Thiếu thư mục", ""],
    ]


def test_r2_latest_done_scan_against_entry_files(db, project):
    rows = build_reconciliation(db, project_id=project.id)["rows"]["R2"]
    by_box = {}
    for row in rows:
        by_box.setdefault(row[0], {})[row[1]] = (row[2], row[3], row[4])
    assert by_box[1] == {
        "012/BIA.pdf": (10, 10, "Khớp"),
        "012/1.PDF": (100, 100, "Khớp"),
        "012/2.pdf": (999, 50, "Khác dung lượng"),
        "012/3.pdf": (70, None, "Thiếu ở nhập liệu"),
        "013a/1.pdf": (60, 60, "Khớp"),
        "014/1.pdf": (5, None, "Thiếu ở nhập liệu"),
        "099/1.pdf": (None, 1, "Thừa ở nhập liệu"),
        "abc/1.pdf": (None, 1, "Thừa ở nhập liệu"),
        "loose.pdf": (None, 1, "Thừa ở nhập liệu"),
    }, "gói done mới nhất (v2); gói cũ v1 và gói failed v3 bỏ qua; file .txt không tính"
    assert by_box[2] == {"001/1.pdf": (200, 200, "Khớp")}
    assert by_box[3] == {"": (None, None, "Hộp chưa có gói scan")}
    assert by_box[4] == {"001/1.pdf": (1, None, "Thiếu ở nhập liệu")}


def test_r3_entry_files_against_submissions(db, project):
    rows = {row[1]: row[2:] for row in build_reconciliation(db, project_id=project.id)["rows"]["R3"]}
    assert rows == {
        "0001/012/1.pdf": [1, "Hoàn thành", "Đạt"],
        "0001/012/2.pdf": [1, "Chờ duyệt", "Chưa hoàn thành (Chờ duyệt)"],
        "0001/013A/1.pdf": [2, "Hoàn thành, Hoàn thành", "Nhiều hồ sơ (2)"],
        "0001/099/1.pdf": [0, "Chưa nhập", "Chưa nhập"],
        "0001/abc/1.pdf": [0, "Chưa nhập", "Chưa nhập"],
        "0001/loose.pdf": [0, "Chưa nhập", "Chưa nhập"],
        "0002/001/1.pdf": [1, "Hoàn thành", "Đạt"],
        "0003/001/1.pdf": [0, "Chưa nhập", "Chưa nhập"],
    }, "bìa và file đã bị thay thế không vào R3"


def test_r4_pages_against_sheets(db, project):
    rows = build_reconciliation(db, project_id=project.id)["rows"]["R4"]
    assert rows == [
        [1, 3, 18, 38, 39, "Hợp lý"],  # 1+8+12+6+11, file lỗi (-1) không tính; ngưỡng 2*18+3
        [2, 1, 4, 20, 9, "Nhiều hơn 2 lần số tờ"],
        [3, 0, None, None, None, "Thiếu dữ liệu"],
        [4, 1, 50, 10, 101, "Ít trang hơn số tờ"],
    ]


def test_summary_and_workbook(db, project):
    report = build_reconciliation(db, project_id=project.id)
    assert report["summary"]["R1"]["results"] == {"Khớp": 3, "Thiếu thư mục": 2, "Thừa thư mục": 3}
    assert report["summary"]["R4"]["results"] == {"Hợp lý": 1, "Nhiều hơn 2 lần số tờ": 1, "Thiếu dữ liệu": 1, "Ít trang hơn số tờ": 1}
    workbook = openpyxl.load_workbook(BytesIO(reconciliation_workbook(report)))
    assert workbook.sheetnames == ["Tổng hợp", "R1 Mục lục-Thư mục", "R2 Scan-Nhập liệu", "R3 File-Hồ sơ nhập", "R4 Trang-Tờ"]
    r4 = workbook["R4 Trang-Tờ"]
    assert [cell.value for cell in r4[1]] == ["Hộp", "Số hồ sơ mục lục", "Tổng số tờ", "Số trang scan", "Ngưỡng trên (2 × tờ + hồ sơ)", "Kết quả"]
    assert r4["F2"].fill.fgColor.rgb.endswith("FFFFFF") or r4["F2"].fill.fill_type is None, "dòng Hợp lý không tô màu"
    assert r4["F3"].fill.fgColor.rgb.endswith("FCE4D6"), "dòng lệch tô màu"
    assert [cell.value for cell in workbook["R3 File-Hồ sơ nhập"][1]] == ["Hộp", "Đường dẫn file", "Số hồ sơ nhập", "Trạng thái", "Kết quả"]
    summary = [[cell.value for cell in row] for row in workbook["Tổng hợp"].iter_rows(min_row=5)]
    assert ["R2", "Gói scan ↔ file nhập liệu", "Hộp chưa có gói scan", 1] in summary


def test_api_json_xlsx_and_permissions(db, project):
    client = TestClient(app)
    app.dependency_overrides[get_db] = lambda: db
    try:
        app.dependency_overrides[get_current_user] = lambda: STAFF
        assert client.get(f"/api/projects/{project.id}/reconciliation").status_code == 403
        app.dependency_overrides[get_current_user] = lambda: ADMIN
        data = client.get(f"/api/projects/{project.id}/reconciliation?format=json").json()["data"]
        assert set(data) == {"project", "summary"}
        assert data["summary"]["R3"]["results"]["Đạt"] == 2
        response = client.get(f"/api/projects/{project.id}/reconciliation")
        assert response.headers["content-disposition"] == 'attachment; filename="Doi_soat_Bo_Y_te.xlsx"'
        assert openpyxl.load_workbook(BytesIO(response.content)).sheetnames[0] == "Tổng hợp"
        missing = client.get("/api/projects/999/reconciliation?format=json")
        assert missing.status_code == 404 and missing.json()["detail"]["code"] == "project_not_found"
    finally:
        app.dependency_overrides.clear()
