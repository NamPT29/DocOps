"""T1: sheet "Chấm công theo ngày" trong file Excel xuất."""

from datetime import date, datetime
from types import SimpleNamespace

import openpyxl
import pytest
from fastapi import BackgroundTasks
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from server import export_worker
from server.models import (
    Base,
    Submission,
    SubmissionQualityAssessment,
    SubmissionReviewHistory,
    Template,
    User,
)
from server.repositories import timesheet_repository
from server.routers import submissions as submissions_router
from server.services import export_job_service
from server.services.excel_service import export_submissions_to_excel
from server.services.timesheet_service import (
    TimesheetRow,
    daily_timesheet_rows,
    vietnam_date,
)


@pytest.fixture()
def db(tmp_path):
    engine = create_engine(f"sqlite:///{(tmp_path / 'timesheet.sqlite3').as_posix()}")
    Base.metadata.create_all(bind=engine)
    session = sessionmaker(bind=engine)()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


@pytest.fixture(autouse=True)
def isolated_heavy_api_limiter(db, monkeypatch):
    from server.services import api_rate_limit_service

    limiter = api_rate_limit_service.DatabaseRateLimiter(
        240, 60, session_factory=sessionmaker(bind=db.get_bind()),
    )
    monkeypatch.setattr(api_rate_limit_service, "heavy_api_rate_limiter", limiter)


def add_user(db, username, full_name=""):
    user = User(username=username, password="hash", role="user", full_name=full_name)
    db.add(user)
    db.flush()
    return user


def add_submission(db, template, creator, *, created_at, status="completed"):
    submission = Submission(
        template_id=template.id,
        data_json='{"col_0": "x"}',
        status=status,
        created_by_user_id=creator.id if creator else None,
        created_at=created_at,
    )
    db.add(submission)
    db.flush()
    return submission


def add_assessment(db, submission, input_user, *, created_at):
    db.add(SubmissionQualityAssessment(
        submission_id=submission.id,
        input_user_id=input_user.id,
        baseline_data_json="{}",
        created_at=created_at,
        updated_at=created_at,
    ))
    db.flush()


def add_review(db, submission, reviewer, *, created_at, event_type="review_confirmed"):
    db.add(SubmissionReviewHistory(
        submission_id=submission.id,
        event_type=event_type,
        input_user_id=submission.created_by_user_id or reviewer.id,
        reviewer_user_id=reviewer.id,
        baseline_data_json="{}",
        reviewer_data_json="{}",
        created_at=created_at,
    ))
    db.flush()


@pytest.fixture()
def scenario(db):
    """Hai người nhập, hai người duyệt, giờ UTC cắt qua nửa đêm giờ Việt Nam."""
    template = Template(name="Mẫu chấm công", filename="timesheet.xlsx")
    db.add(template)
    db.flush()
    an = add_user(db, "an", "Nguyễn An")
    binh = add_user(db, "binh", "Trần Bình")
    chi = add_user(db, "chi")  # chưa có họ tên -> dùng tên đăng nhập
    dung = add_user(db, "dung", "Lê Dũng")

    # 06/10 17:30 UTC = 07/10 00:30 giờ VN.
    s1 = add_submission(db, template, an, created_at=datetime(2026, 10, 9, 1, 0))
    add_assessment(db, s1, an, created_at=datetime(2026, 10, 6, 17, 30))
    # 07/10 16:59 UTC = 07/10 23:59 giờ VN.
    s2 = add_submission(db, template, an, created_at=datetime(2026, 10, 9, 1, 0))
    add_assessment(db, s2, an, created_at=datetime(2026, 10, 7, 16, 59))
    # 07/10 17:00 UTC = 08/10 00:00 giờ VN.
    s3 = add_submission(db, template, binh, created_at=datetime(2026, 10, 9, 1, 0))
    add_assessment(db, s3, binh, created_at=datetime(2026, 10, 7, 17, 0))
    # Không có baseline: lấy người tạo và created_at của hồ sơ.
    s4 = add_submission(db, template, binh, created_at=datetime(2026, 10, 7, 3, 0))
    # Baseline ghi người nhập khác người tạo hồ sơ: theo baseline.
    s5 = add_submission(db, template, dung, created_at=datetime(2026, 10, 9, 1, 0))
    add_assessment(db, s5, an, created_at=datetime(2026, 10, 7, 2, 0))

    # Chỉ lần review_confirmed ĐẦU TIÊN được tính.
    add_review(db, s1, chi, created_at=datetime(2026, 10, 7, 9, 0))
    add_review(db, s1, dung, created_at=datetime(2026, 10, 8, 9, 0))
    add_review(db, s2, chi, created_at=datetime(2026, 10, 7, 17, 30))  # 08/10 giờ VN
    # Xác nhận chỉnh sửa của người nhập không phải lần duyệt.
    add_review(db, s3, dung, created_at=datetime(2026, 10, 8, 2, 0), event_type="input_confirmed")

    # Hồ sơ không có trong file xuất: không được tính.
    outside = add_submission(db, template, an, created_at=datetime(2026, 10, 7, 3, 0))
    add_assessment(db, outside, an, created_at=datetime(2026, 10, 7, 3, 0))
    add_review(db, outside, chi, created_at=datetime(2026, 10, 7, 4, 0))
    db.commit()
    return SimpleNamespace(template=template, exported=[s1, s2, s3, s4, s5], outside=outside)


EXPECTED_ROWS = [
    TimesheetRow(date(2026, 10, 7), "chi", entered=0, reviewed=1),
    TimesheetRow(date(2026, 10, 7), "Nguyễn An", entered=3, reviewed=0),
    TimesheetRow(date(2026, 10, 7), "Trần Bình", entered=1, reviewed=0),
    TimesheetRow(date(2026, 10, 8), "chi", entered=0, reviewed=1),
    TimesheetRow(date(2026, 10, 8), "Trần Bình", entered=1, reviewed=0),
]


def test_vietnam_date_converts_naive_utc():
    assert vietnam_date(datetime(2026, 10, 6, 16, 59)) == date(2026, 10, 6)
    assert vietnam_date(datetime(2026, 10, 6, 17, 0)) == date(2026, 10, 7)
    assert vietnam_date(None) is None


def test_daily_rows_count_entered_and_first_review_per_person_per_vietnam_day(db, scenario):
    assert daily_timesheet_rows(db, scenario.exported) == EXPECTED_ROWS


def test_daily_rows_only_cover_exported_submissions(db, scenario):
    rows = daily_timesheet_rows(db, [scenario.outside])

    assert rows == [
        TimesheetRow(date(2026, 10, 7), "chi", entered=0, reviewed=1),
        TimesheetRow(date(2026, 10, 7), "Nguyễn An", entered=1, reviewed=0),
    ]


def test_daily_rows_empty_export():
    assert daily_timesheet_rows(None, []) == []


def test_large_exports_are_queried_in_chunks(db, scenario, monkeypatch):
    monkeypatch.setattr(timesheet_repository, "ID_CHUNK_SIZE", 2)

    assert daily_timesheet_rows(db, scenario.exported) == EXPECTED_ROWS


def _template(tmp_path, *extra_sheets):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"
    worksheet.append(["Nhóm", None])
    worksheet.append(["Cột A", "Cột B"])
    worksheet.append(["Nhãn A", "Nhãn B"])
    for title in extra_sheets:
        workbook.create_sheet(title).append(["giữ nguyên"])
    path = tmp_path / "template.xlsx"
    workbook.save(path)
    return path


def test_export_appends_timesheet_sheet_at_the_end(tmp_path):
    template_path = _template(tmp_path, "Ghi chú")
    output_path = tmp_path / "out.xlsx"
    submissions = [SimpleNamespace(data_json='{"col_0": "A1", "col_1": "B1"}')]

    export_submissions_to_excel(template_path, submissions, output_path, timesheet_rows=EXPECTED_ROWS)

    workbook = openpyxl.load_workbook(output_path)
    try:
        assert workbook.sheetnames == ["Data", "Ghi chú", "Chấm công theo ngày"]
        assert workbook.active.title == "Data"
        assert [workbook["Data"].cell(4, column).value for column in (1, 2)] == ["A1", "B1"]
        sheet = workbook["Chấm công theo ngày"]
        values = [list(row) for row in sheet.iter_rows(values_only=True)]
        assert values[0] == ["Ngày", "Họ và tên", "Số hàng đã nhập", "Số hàng đã duyệt"]
        assert values[1:] == [
            [datetime(row.work_date.year, row.work_date.month, row.work_date.day),
             row.name, row.entered, row.reviewed]
            for row in EXPECTED_ROWS
        ]
        assert sheet.cell(2, 1).number_format == "DD/MM/YYYY"
        assert sheet.freeze_panes == "A2"
    finally:
        workbook.close()


def test_export_without_timesheet_rows_keeps_template_sheets(tmp_path):
    template_path = _template(tmp_path)
    output_path = tmp_path / "out.xlsx"

    export_submissions_to_excel(template_path, [], output_path)

    workbook = openpyxl.load_workbook(output_path)
    try:
        assert workbook.sheetnames == ["Data"]
    finally:
        workbook.close()


def test_export_with_no_rows_still_has_header_and_avoids_name_clash(tmp_path):
    template_path = _template(tmp_path, "Chấm công theo ngày")
    output_path = tmp_path / "out.xlsx"

    export_submissions_to_excel(template_path, [], output_path, timesheet_rows=[])

    workbook = openpyxl.load_workbook(output_path)
    try:
        assert workbook.sheetnames == ["Data", "Chấm công theo ngày", "Chấm công theo ngày (2)"]
        assert workbook["Chấm công theo ngày"]["A1"].value == "giữ nguyên"
        assert workbook["Chấm công theo ngày (2)"].max_row == 1
    finally:
        workbook.close()


def test_export_endpoint_passes_timesheet_of_exported_submissions(db, scenario, mocker):
    export_call = mocker.patch("server.services.excel_service.export_submissions_to_excel")
    mocker.patch("server.routers.submissions.FileResponse", return_value={"status": "file-ready"})

    result = submissions_router.api_export(
        scenario.template.id,
        BackgroundTasks(),
        current_user={"id": 1, "role": "admin"},
        db=db,
    )

    assert result == {"status": "file-ready"}
    exported_ids = [submission.id for submission in export_call.call_args.args[1]]
    assert sorted(exported_ids) == sorted(
        [submission.id for submission in scenario.exported] + [scenario.outside.id]
    )
    rows = export_call.call_args.kwargs["timesheet_rows"]
    assert sum(row.entered for row in rows) == 6
    assert sum(row.reviewed for row in rows) == 3


def test_export_worker_passes_timesheet_rows(db, scenario, tmp_path, monkeypatch):
    monkeypatch.setattr(export_job_service, "EXPORT_SCRATCH_DIR", tmp_path)
    monkeypatch.setattr(export_job_service, "EXPORT_LOCK_PATH", tmp_path / "export_all.lock")
    job_id = "f" * 32
    export_job_service.write_export_job(job_id, {
        "job_id": job_id, "state": "queued", "extension": ".xlsx", "filename": "t.xlsx",
    })
    export_job_service.EXPORT_LOCK_PATH.write_text(job_id, encoding="utf-8")
    monkeypatch.setattr(export_worker, "SessionLocal", sessionmaker(bind=db.get_bind()))
    captured = {}

    def fake_export(_template_path, exported, output_path, **kwargs):
        captured["ids"] = sorted(submission.id for submission in exported)
        captured["rows"] = kwargs.get("timesheet_rows")
        open(output_path, "wb").close()

    monkeypatch.setattr(export_worker, "export_submissions_to_excel", fake_export)
    args = SimpleNamespace(
        job_id=job_id, extension=".xlsx", project_id=None, template_id=scenario.template.id,
        include_pending_review=False, folder_path=None, start_date=None, end_date=None,
    )

    assert export_worker.run_export_job(args) == 0
    assert captured["rows"] == daily_timesheet_rows(db, [
        submission for submission in scenario.exported + [scenario.outside]
    ])
    assert export_job_service.read_export_job(job_id)["state"] == "completed"
