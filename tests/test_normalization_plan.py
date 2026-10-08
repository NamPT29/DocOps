"""Kế hoạch chuẩn hóa (lát G1): mã hồ sơ, mã văn bản, đường dẫn bàn giao QC-03/QC-04. Chỉ đọc."""
import json
from io import BytesIO

import openpyxl
import pytest
from fastapi import HTTPException
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from server.database import Base
from server.models import (
    ArrangementDossier,
    ArrangementImport,
    AssignedDocument,
    Project,
    ProjectCase,
    ProjectDocumentAsset,
    ProjectPolicy,
    ProjectReportUnit,
    Submission,
    Template,
    User,
)
from server.routers import projects
from server.services import normalization_plan_service as plan_service
from server.services.normalization_plan_service import (
    PROBLEM_BAD_FOLDER,
    PROBLEM_MAINTENANCE,
    PROBLEM_NO_ORGAN,
    PROBLEM_NOT_APPROVED,
    PROBLEM_NOT_IN_CATALOG,
    PROBLEM_NOT_IN_DOSSIER,
    ascii_name,
    build_plan,
    dossier_code,
    is_cover_file,
    natural_key,
    plan_workbook,
)

ROOT = "CSDL_SOHOA_Bo_Y_te_2026"
ADMIN = {"id": 1, "username": "admin", "full_name": "Admin", "role": "admin", "session_id": "s"}


@pytest.fixture()
def db(tmp_path):
    engine = create_engine(f"sqlite:///{(tmp_path / 'plan.sqlite3').as_posix()}")
    Base.metadata.create_all(bind=engine)
    session = sessionmaker(bind=engine)()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


def dossier(project, case_row, number, *, suffix="", maintenance="01", notation=None, year=2006):
    return ArrangementDossier(
        project_id=project.id, case_id=case_row.id, box_number=1, dossier_number=number,
        dossier_suffix=suffix, fonds_code="Phông BYT", fonds_name="Bộ Y tế", catalog_number="01",
        file_notation=notation, title=f"Hồ sơ {number}", start_date=f"01/01/{year}",
        end_date=f"31/12/{year}", start_year=year, maintenance_code=maintenance, sheet_count=10,
        source_row=number,
    )


@pytest.fixture()
def world(db):
    admin = User(username="admin", password="x", role="admin")
    template = Template(name="T", filename="t.xlsx")
    db.add_all([admin, template])
    db.flush()
    project = Project(
        name="Bộ Y tế 2026", root_folder_name="Goc", template_id=template.id,
        template_name_snapshot="T", template_filename_snapshot="t.xlsx", case_level=1,
        report_mode="pdf", created_by_user_id=admin.id, status="new",
    )
    db.add(project)
    db.flush()
    case_row = ProjectCase(project_id=project.id, case_key="0001", display_name="0001")
    db.add(case_row)
    db.flush()
    report = ProjectReportUnit(project_id=project.id, case_id=case_row.id, report_key="0001", display_name="0001")
    policy = ProjectPolicy(project_id=project.id, organ_code="H05.02.02")
    db.add_all([
        report, policy,
        dossier(project, case_row, 12),
        dossier(project, case_row, 13, suffix="a", maintenance="02", notation="HC"),
        dossier(project, case_row, 14, maintenance="09"),
    ])
    db.flush()
    later_import = ArrangementImport(project_id=project.id, file_name="muc_luc_2.xlsx", file_sha256="1" * 64)
    db.add(later_import)
    db.flush()
    dropped = dossier(project, case_row, 15)
    dropped.missing_from_import_id = later_import.id
    db.add(dropped)
    db.flush()

    def asset(path, *statuses, asset_status="active"):
        document = None
        if statuses:
            document = AssignedDocument(original_filename=path, uuid_filename=f"u-{path}")
            db.add(document)
            db.flush()
            for status in statuses:  # lần lưu sau có id lớn hơn
                db.add(Submission(template_id=template.id, assigned_document_id=document.id, data_json="{}", status=status))
                db.flush()
        db.add(ProjectDocumentAsset(
            project_id=project.id, case_id=case_row.id, report_unit_id=report.id,
            assigned_document_id=document.id if document else None, relative_path=path,
            normalized_relative_path=path.casefold(), original_filename=path.rsplit("/", 1)[-1],
            storage_filename=f"s-{path.replace('/', '-')}-{asset_status}", content_sha256="0" * 64,
            byte_size=1, status=asset_status,
        ))

    asset("0001/012/BIA.pdf")
    asset("0001/012/10.pdf", "pending_review")
    asset("0001/012/2.pdf", "draft", "completed")
    asset("0001/012/3.pdf", "completed", asset_status="replaced")
    asset("0001/015/1.pdf", "completed")
    asset("0001/013a/1.pdf", "completed")
    asset("0001/014/1.pdf", "completed")
    asset("0001/099/1.pdf", "completed")
    asset("0001/abc/1.pdf", "completed")
    asset("0001/loose.pdf", "completed")
    asset("0001/012/con/5.pdf", "completed")
    db.commit()
    return project, policy


def rows_by_path(plan):
    return {row["relative_path"]: row for row in plan["rows"]}


def test_codes_and_paths_follow_qc03_qc04(db, world):
    project, _policy = world
    rows = rows_by_path(build_plan(db, project_id=project.id))

    base = f"{ROOT}/H05.02.02/2006/VV/Phong_BYT/H05.02.02.2006.12"
    assert rows["0001/012/2.pdf"] == {
        "box": 1, "folder": "012", "relative_path": "0001/012/2.pdf", "kind": "Văn bản", "number": 1,
        "dossier_code": "H05.02.02.2006.12", "document_code": "H05.02.02.2006.12.0000001",
        "target_path": f"{base}/H05.02.02.2006.12.0000001.pdf", "entry_status": "Hoàn thành", "problems": [],
    }
    assert rows["0001/012/10.pdf"]["number"] == 2, "STT theo thứ tự tự nhiên: 2.pdf trước 10.pdf"
    assert rows["0001/012/10.pdf"]["target_path"] == f"{base}/H05.02.02.2006.12.0000002.pdf"
    cover = rows["0001/012/BIA.pdf"]
    assert (cover["kind"], cover["number"], cover["document_code"]) == ("Bìa", None, "")
    assert cover["target_path"] == f"{base}/H05.02.02.2006.12_BIA.pdf"
    assert cover["problems"] == [], "file bìa không cần nhập liệu"
    assert "0001/012/3.pdf" not in rows, "file đã bị thay thế không vào kế hoạch"


def test_suffix_notation_and_long_term_maintenance(db, world):
    project, _policy = world
    row = rows_by_path(build_plan(db, project_id=project.id))["0001/013a/1.pdf"]

    assert row["dossier_code"] == "H05.02.02.2006.13a.HC"
    assert row["target_path"] == f"{ROOT}/H05.02.02/2006/LD/Phong_BYT/H05.02.02.2006.13a.HC/H05.02.02.2006.13a.HC.0000001.pdf"


def test_project_file_notation_is_the_default(db, world):
    project, policy = world
    policy.file_notation = "TC"
    db.commit()
    rows = rows_by_path(build_plan(db, project_id=project.id))

    assert rows["0001/012/2.pdf"]["dossier_code"] == "H05.02.02.2006.12.TC"
    assert rows["0001/013a/1.pdf"]["dossier_code"] == "H05.02.02.2006.13a.HC", "ký hiệu của mục lục thắng"


def test_every_problem_is_reported_per_file(db, world):
    project, _policy = world
    rows = rows_by_path(build_plan(db, project_id=project.id))

    assert rows["0001/012/10.pdf"]["problems"] == [PROBLEM_NOT_APPROVED]
    assert rows["0001/012/10.pdf"]["entry_status"] == "Chờ duyệt"
    assert rows["0001/014/1.pdf"]["problems"] == [PROBLEM_MAINTENANCE]
    assert rows["0001/014/1.pdf"]["dossier_code"] == "H05.02.02.2006.14"
    assert rows["0001/014/1.pdf"]["target_path"] == "", "không có viết tắt THBQ thì không dựng được đường dẫn"
    assert rows["0001/099/1.pdf"]["problems"] == [PROBLEM_NOT_IN_CATALOG]
    assert rows["0001/015/1.pdf"]["problems"] == [PROBLEM_NOT_IN_CATALOG], "dòng mục lục đã bị loại ở lần nhập sau"
    assert rows["0001/012/10.pdf"]["entry_status"] == "Chờ duyệt"
    assert rows["0001/abc/1.pdf"]["problems"] == [PROBLEM_BAD_FOLDER]
    assert rows["0001/loose.pdf"]["problems"] == [PROBLEM_NOT_IN_DOSSIER]
    assert rows["0001/012/con/5.pdf"]["problems"] == [PROBLEM_NOT_IN_DOSSIER], "thư mục con trong hồ sơ là sai cấu trúc"
    for path in ("0001/099/1.pdf", "0001/015/1.pdf", "0001/abc/1.pdf", "0001/loose.pdf", "0001/012/con/5.pdf"):
        assert rows[path]["dossier_code"] == rows[path]["target_path"] == ""


def test_missing_organ_code_blocks_every_file(db, world):
    project, policy = world
    policy.organ_code = None
    db.commit()
    plan = build_plan(db, project_id=project.id)

    assert all(PROBLEM_NO_ORGAN in row["problems"] for row in plan["rows"])
    assert all(row["dossier_code"] == "" and row["target_path"] == "" for row in plan["rows"])
    assert plan["summary"]["ready"] == 0


def test_summary_counts(db, world):
    project, _policy = world
    summary = build_plan(db, project_id=project.id)["summary"]

    assert summary == {
        "files": 10, "documents": 9, "covers": 1, "dossiers": 3, "ready": 3, "with_problems": 7,
        "problems": {
            PROBLEM_NOT_APPROVED: 1, PROBLEM_MAINTENANCE: 1, PROBLEM_NOT_IN_CATALOG: 2,
            PROBLEM_BAD_FOLDER: 1, PROBLEM_NOT_IN_DOSSIER: 2,
        },
    }


def test_workbook_has_summary_and_plan_sheets(db, world):
    project, _policy = world
    workbook = openpyxl.load_workbook(BytesIO(plan_workbook(build_plan(db, project_id=project.id))))

    assert workbook.sheetnames == ["Tổng hợp", "Kế hoạch đổi tên"]
    summary = {row[0]: row[1] for row in workbook["Tổng hợp"].iter_rows(values_only=True)}
    assert summary["Mã cơ quan"] == "H05.02.02"
    assert summary["Thư mục gốc bàn giao"] == ROOT
    assert summary["Số file còn vấn đề"] == 7
    sheet = workbook["Kế hoạch đổi tên"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][:3] == ("Hộp", "Thư mục hồ sơ", "File hiện tại")
    assert len(rows) == 11
    by_path = {row[2]: row for row in rows[1:]}
    assert by_path["0001/012/10.pdf"][6:] == (
        "H05.02.02.2006.12.0000002",
        f"{ROOT}/H05.02.02/2006/VV/Phong_BYT/H05.02.02.2006.12/H05.02.02.2006.12.0000002.pdf",
        "Chờ duyệt", PROBLEM_NOT_APPROVED,
    )


def test_endpoint_returns_json_summary_and_xlsx(db, world):
    project, _policy = world

    data = projects.api_get_normalization_plan(project.id, format="json", current_user=ADMIN, db=db)["data"]
    assert data["summary"]["files"] == 10 and data["root"] == ROOT and "rows" not in data

    response = projects.api_get_normalization_plan(project.id, format="xlsx", current_user=ADMIN, db=db)
    assert response.media_type.endswith("spreadsheetml.sheet")
    assert response.headers["content-disposition"] == 'attachment; filename="Ke_hoach_chuan_hoa_Bo_Y_te_2026.xlsx"'
    assert openpyxl.load_workbook(BytesIO(response.body)).sheetnames == ["Tổng hợp", "Kế hoạch đổi tên"]

    with pytest.raises(HTTPException) as error:
        build_plan(db, project_id=999)
    assert error.value.status_code == 404


def test_only_admin_may_call_the_endpoint():
    route = next(r for r in projects.router.routes if r.path.endswith("/normalization-plan"))
    dependencies = {dependency.call for dependency in route.dependant.dependencies}
    assert projects.get_admin_user in dependencies


@pytest.mark.parametrize("text, expected", [
    ("Phông Bộ Y tế / số 1", "Phong_Bo_Y_te__so_1"),
    ("Đảng ủy", "Dang_uy"),
    ("H05.02-a_b", "H05.02-a_b"),
])
def test_ascii_name_keeps_only_qc04_characters(text, expected):
    assert ascii_name(text) == expected


def test_helpers():
    assert sorted(["10.pdf", "2.pdf", "1.pdf", "b.pdf"], key=natural_key) == ["1.pdf", "2.pdf", "10.pdf", "b.pdf"]
    assert is_cover_file("BIA.pdf") and is_cover_file("Bìa hồ sơ.pdf") and not is_cover_file("0000001.pdf")
    assert dossier_code("X", 2008, 5, "", None) == "X.2008.05"
    assert dossier_code("X", 2008, 100, "b", "HC") == "X.2008.100b.HC"
    assert plan_service.MAINTENANCE_ABBREVIATIONS == {"01": "VV", "02": "LD"}
