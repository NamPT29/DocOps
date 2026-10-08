"""Đóng gói bàn giao (lát G2): chép file theo kế hoạch G1, SHA-256, metadata NN-SIP, chạy lại, khóa."""
import hashlib
import json
import os
from io import BytesIO
from pathlib import Path

import openpyxl
import pypdf
import pytest
from fastapi import HTTPException
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from server.database import Base
from server.models import (
    ArrangementDossier,
    AssignedDocument,
    Project,
    ProjectCase,
    ProjectDocumentAsset,
    ProjectPolicy,
    ProjectReportUnit,
    Submission,
    Template,
    User,
)
from server.routers import projects
from server.services import handover_package_service as package
from server.services.normalization_plan_service import PROBLEM_NOT_APPROVED

ROOT = "CSDL_SOHOA_Bo_Y_te_2026"
ADMIN = {"id": 1, "username": "admin", "full_name": "Admin", "role": "admin", "session_id": "s"}
LABELS = [
    "Tiêu đề hồ sơ", "Thời gian bắt đầu", "Thời gian kết thúc", "Số tờ",
    "Tên cơ quan, tổ chức ban hành văn bản", "Số của văn bản", "Ký hiệu của văn bản",
    "Ngày ký (Ngày, tháng, năm văn bản)", "Thể loại văn bản", "Trích yếu nội dung", "Người ký",
]
SCHEMA = [
    {"category": "Thông tin hồ sơ (Bìa)", "fields": [{"col_index": i, "label": LABELS[i], "name": f"col_{i}"} for i in range(4)]},
    {"category": "Thông tin văn bản", "fields": [{"col_index": i, "label": LABELS[i], "name": f"col_{i}"} for i in range(4, 11)]},
]


def pdf_bytes(pages: int) -> bytes:
    writer = pypdf.PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=595, height=842)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


@pytest.fixture()
def env(tmp_path, monkeypatch):
    storage = tmp_path / "kho"
    handover = tmp_path / "ban_giao"
    storage.mkdir()
    monkeypatch.setenv("PDF_STORAGE_PATH", str(storage))
    monkeypatch.setenv("HANDOVER_DIR", str(handover))
    engine = create_engine(f"sqlite:///{(tmp_path / 'package.sqlite3').as_posix()}")
    Base.metadata.create_all(bind=engine)
    factory = sessionmaker(bind=engine)
    db = factory()
    try:
        yield {"db": db, "factory": factory, "storage": storage, "handover": handover}
    finally:
        db.close()
        engine.dispose()


def dossier(project, case_row, number, *, suffix="", notation=None):
    return ArrangementDossier(
        project_id=project.id, case_id=case_row.id, box_number=1, dossier_number=number, dossier_suffix=suffix,
        fonds_code="Phông BYT", fonds_name="Bộ Y tế", catalog_number="01", file_notation=notation,
        title=f"  Tập  quyết định {number}  ", start_date="01/01/2006", end_date="31/12/2006", start_year=2006,
        maintenance_code="01", sheet_count=20 + number, term="2006-2010", note="Ghi chú mục lục", source_row=number,
    )


@pytest.fixture()
def world(env):
    db, storage = env["db"], env["storage"]
    admin = User(username="admin", password="x", role="admin")
    template = Template(name="T", filename="t.xlsx")
    db.add_all([admin, template])
    db.flush()
    project = Project(
        name="Bộ Y tế 2026", root_folder_name="Goc", template_id=template.id, template_name_snapshot="T",
        template_filename_snapshot="t.xlsx", form_schema_json_snapshot=json.dumps(SCHEMA, ensure_ascii=False),
        case_level=1, report_mode="pdf", created_by_user_id=admin.id, status="new",
    )
    db.add(project)
    db.flush()
    case_row = ProjectCase(project_id=project.id, case_key="0001", display_name="0001")
    db.add(case_row)
    db.flush()
    report = ProjectReportUnit(project_id=project.id, case_id=case_row.id, report_key="0001", display_name="0001")
    db.add_all([
        report, ProjectPolicy(project_id=project.id, organ_code="H05.02.02"),
        dossier(project, case_row, 12), dossier(project, case_row, 13, suffix="a", notation="HC"), dossier(project, case_row, 14),
    ])
    db.flush()
    files = {}

    def asset(path, pages, status=None, data=None):
        content = pdf_bytes(pages)
        storage_name = f"s-{len(files)}.pdf"
        (storage / storage_name).write_bytes(content)
        document = None
        if status:
            document = AssignedDocument(original_filename=path, uuid_filename=f"u-{path}")
            db.add(document)
            db.flush()
            db.add(Submission(template_id=template.id, assigned_document_id=document.id, status=status,
                              data_json=json.dumps(data or {}, ensure_ascii=False)))
        db.add(ProjectDocumentAsset(
            project_id=project.id, case_id=case_row.id, report_unit_id=report.id,
            assigned_document_id=document.id if document else None, relative_path=path,
            normalized_relative_path=path.casefold(), original_filename=path.rsplit("/", 1)[-1],
            storage_filename=storage_name, content_sha256=hashlib.sha256(content).hexdigest(), byte_size=len(content),
        ))
        files[path] = {"content": content, "storage": storage / storage_name}

    entry = {
        "col_0": "Tiêu đề trên bìa", "col_4": "  bộ y tế ", "col_5": "3282", "col_6": "/QĐ-BYT", "col_7": "01/09/2006",
        "col_8": "Quyết định", "col_9": "Phê duyệt   báo cáo\nkinh tế", "col_10": "Trần Chí Liêm",
    }
    asset("0001/012/BIA.pdf", 1)
    asset("0001/012/10.pdf", 3, "completed", {**entry, "col_5": "3283"})
    asset("0001/012/2.pdf", 2, "completed", entry)
    asset("0001/013a/1.pdf", 4, "completed", {**entry, "col_8": "Công văn"})
    asset("0001/014/1.pdf", 1, "pending_review", entry)
    db.commit()
    return {**env, "project": project, "files": files}


def run(world):
    return package.start_package(
        world["db"], project_id=world["project"].id, current_user=ADMIN,
        session_factory=world["factory"], runner=package.run_package,
    )


def target(world, relative):
    return world["handover"] / ROOT / "H05.02.02/2006/VV/Phong_BYT" / relative


def sheet_rows(path, name):
    rows = list(openpyxl.load_workbook(path)[name].iter_rows(values_only=True))
    return [dict(zip(rows[0], row)) for row in rows[1:]]


def test_ready_dossiers_are_copied_with_handover_names(world):
    queued = run(world)
    job = package.read_job(world["project"].id)

    assert queued["state"] == "queued" and queued["files_total"] == 4
    assert job["state"] == "done", job["message"]
    assert (job["dossiers_packaged"], job["files_copied"], job["files_skipped"], job["error_count"]) == (2, 4, 0, 0)
    assert job["dossiers_skipped"] == 1
    base = "H05.02.02.2006.12"
    for relative, source in (
        (f"{base}/{base}_BIA.pdf", "0001/012/BIA.pdf"),
        (f"{base}/{base}.0000001.pdf", "0001/012/2.pdf"),
        (f"{base}/{base}.0000002.pdf", "0001/012/10.pdf"),
        ("H05.02.02.2006.13a.HC/H05.02.02.2006.13a.HC.0000001.pdf", "0001/013a/1.pdf"),
    ):
        assert target(world, relative).read_bytes() == world["files"][source]["content"]
    assert not target(world, "H05.02.02.2006.14").exists(), "hồ sơ còn văn bản chưa duyệt thì không đóng gói"
    assert not list(world["handover"].rglob("*.part"))
    assert not (world["handover"] / "_jobs" / f"project_{world['project'].id}.lock").exists(), "xong phải nhả khóa"


def test_sha256sums_cover_every_packaged_file(world):
    run(world)
    lines = (world["handover"] / ROOT / "SHA256SUMS.txt").read_text(encoding="utf-8").splitlines()

    assert len(lines) == 4
    for line in lines:
        digest, relative = line.split("  ", 1)
        assert hashlib.sha256((world["handover"] / ROOT / relative).read_bytes()).hexdigest() == digest
        assert not relative.startswith(ROOT)


def test_metadata_nn_sip_sheets(world):
    run(world)
    path = world["handover"] / ROOT / "Metadata_NN-SIP.xlsx"

    assert openpyxl.load_workbook(path).sheetnames == ["Metadata_HS", "MetadataVB"]
    hs = {row["fileCode"]: row for row in sheet_rows(path, "Metadata_HS")}
    assert set(hs) == {"H05.02.02.2006.12", "H05.02.02.2006.13a.HC"}
    first = hs["H05.02.02.2006.12"]
    assert list(first)[:18] == list(package.HS_FIELDS)
    assert (first["title"], first["maintenance"], first["mode"], first["language"]) == ("Tập quyết định 12", "01", "01", "01")
    assert (first["startDate"], first["endDate"], first["numberOfPaper"]) == ("01/01/2006", "31/12/2006", 32)
    assert (first["totalDoc"], first["numberOfPage"], first["confidenceLevel"], first["format"]) == (2, 6, "02", "Bình thường")
    assert first["Tệp tin hồ sơ"] == "H05.02.02.2006.12_BIA.pdf"
    assert (first["Hộp số"], first["Hồ sơ số"], first["Mã phông"], first["Giai đoạn/Nhiệm kỳ"]) == (1, "12", "Phông BYT", "2006-2010")
    assert first["Path"] == f"{ROOT}/H05.02.02/2006/VV/Phong_BYT/H05.02.02.2006.12"
    assert hs["H05.02.02.2006.13a.HC"]["Hồ sơ số"] == "13a"

    vb = {row["docCode"]: row for row in sheet_rows(path, "MetadataVB")}
    assert set(vb) == {"H05.02.02.2006.12.0000001", "H05.02.02.2006.12.0000002", "H05.02.02.2006.13a.HC.0000001"}
    doc = vb["H05.02.02.2006.12.0000001"]
    assert list(doc)[:22] == list(package.VB_FIELDS)
    assert (doc["fileCode"], doc["docId"], doc["typeName"], doc["organName"]) == ("H05.02.02.2006.12", 1, "QUYẾT ĐỊNH", "BỘ Y TẾ")
    assert (doc["codeNumber"], doc["codeNotation"], doc["issuedDate"]) == ("3282", "/QĐ-BYT", "01/09/2006")
    assert doc["subject"] == "Phê duyệt báo cáo kinh tế", "QC-15: gộp khoảng trắng, bỏ xuống dòng"
    assert (doc["numberOfPage"], doc["Người ký"], doc["File gốc"]) == (2, "Trần Chí Liêm", "0001/012/2.pdf")
    assert vb["H05.02.02.2006.12.0000002"]["codeNumber"] == "3283"
    assert vb["H05.02.02.2006.13a.HC.0000001"]["typeName"] == "CÔNG VĂN"


def test_log_lists_packaged_and_skipped_dossiers(world):
    run(world)
    path = world["handover"] / ROOT / "Nhat_ky_dong_goi.xlsx"

    done = sheet_rows(path, "Đã đóng gói")
    assert len(done) == 4 and {row["Kết quả"] for row in done} == {"Đã chép"}
    skipped = sheet_rows(path, "Chưa đóng gói")
    assert [(row["Thư mục hồ sơ"], row["Lý do"]) for row in skipped] == [("014", PROBLEM_NOT_APPROVED)]
    assert sheet_rows(path, "Lỗi chép file") == []


def test_second_run_skips_files_that_are_already_correct(world):
    run(world)
    job = run(world) and package.read_job(world["project"].id)

    assert (job["state"], job["files_copied"], job["files_skipped"]) == ("done", 0, 4)
    log = sheet_rows(world["handover"] / ROOT / "Nhat_ky_dong_goi.xlsx", "Đã đóng gói")
    assert {row["Kết quả"] for row in log} == {"Đã có sẵn, đúng SHA-256"}


def test_damaged_or_missing_source_is_not_packaged(world):
    world["files"]["0001/012/10.pdf"]["storage"].write_bytes(b"hong")
    world["files"]["0001/013a/1.pdf"]["storage"].unlink()
    run(world)
    job = package.read_job(world["project"].id)

    assert (job["state"], job["dossiers_packaged"], job["error_count"]) == ("done", 0, 2)
    assert not target(world, "H05.02.02.2006.12/H05.02.02.2006.12.0000002.pdf").exists()
    assert not list(world["handover"].rglob("*.part"))
    path = world["handover"] / ROOT / "Nhat_ky_dong_goi.xlsx"
    errors = {row["File hiện tại"]: row["Lỗi"] for row in sheet_rows(path, "Lỗi chép file")}
    assert "SHA-256" in errors["0001/012/10.pdf"] and "Không tìm thấy" in errors["0001/013a/1.pdf"]
    assert sheet_rows(world["handover"] / ROOT / "Metadata_NN-SIP.xlsx", "Metadata_HS") == []
    reasons = {row["Thư mục hồ sơ"]: row["Lý do"] for row in sheet_rows(path, "Chưa đóng gói")}
    assert reasons["012"].startswith("Lỗi chép file") and reasons["013a"].startswith("Lỗi chép file")


def test_duplicate_dossier_codes_are_both_skipped(world):
    db = world["db"]
    original = db.query(ProjectDocumentAsset).filter_by(relative_path="0001/013a/1.pdf").one()
    content = pdf_bytes(1)
    (world["storage"] / "dup.pdf").write_bytes(content)
    db.add(ProjectDocumentAsset(
        project_id=original.project_id, case_id=original.case_id, report_unit_id=original.report_unit_id,
        relative_path="0001/0012/1.pdf", normalized_relative_path="0001/0012/1.pdf", original_filename="1.pdf",
        storage_filename="dup.pdf", content_sha256=hashlib.sha256(content).hexdigest(), byte_size=len(content),
    ))
    db.commit()
    context = package.collect_package(db, world["project"].id)

    skipped = {item["folder"]: item["reasons"] for item in context["skipped"]}
    assert "Trùng mã hồ sơ với thư mục khác" in skipped["012"] and "Trùng mã hồ sơ với thư mục khác" in skipped["0012"]
    assert [items[0]["folder"] for items in context["ready"]] == ["013a"]


def test_start_checks_project_ready_dossiers_and_lock(world):
    db, project_id = world["db"], world["project"].id
    with pytest.raises(HTTPException) as error:
        package.start_package(db, project_id=999, current_user=ADMIN, session_factory=world["factory"], runner=lambda *a: None)
    assert error.value.status_code == 404

    policy = db.query(ProjectPolicy).one()
    policy.organ_code = None
    db.commit()
    with pytest.raises(HTTPException) as error:
        run(world)
    assert error.value.status_code == 409 and "Kế hoạch chuẩn hóa" in error.value.detail
    assert not package._lock_path(project_id).exists(), "báo lỗi thì phải nhả khóa"

    policy.organ_code = "H05.02.02"
    db.commit()
    package._lock_path(project_id).parent.mkdir(parents=True, exist_ok=True)
    package._lock_path(project_id).write_text(str(os.getpid()), encoding="utf-8")
    with pytest.raises(HTTPException) as error:
        run(world)
    assert error.value.status_code == 409 and "đang được đóng gói" in error.value.detail

    package._lock_path(project_id).write_text("999999999", encoding="utf-8")
    assert run(world)["state"] == "queued", "khóa của tiến trình đã chết thì được lấy lại"
    assert package.read_job(project_id)["state"] == "done"


def test_unexpected_failure_is_reported_and_releases_lock(world, monkeypatch):
    monkeypatch.setattr(package, "_write_outputs", lambda *args: (_ for _ in ()).throw(RuntimeError("đĩa đầy")))
    run(world)
    job = package.read_job(world["project"].id)

    assert job["state"] == "error" and "đĩa đầy" in job["message"]
    assert not package._lock_path(world["project"].id).exists()


def test_startup_marks_jobs_of_dead_processes_failed(env):
    jobs = env["handover"] / "_jobs"
    jobs.mkdir(parents=True)
    (jobs / "project_7.json").write_text(json.dumps({"project_id": 7, "state": "running", "pid": 999999999}), encoding="utf-8")
    (jobs / "project_7.lock").write_text("999999999", encoding="utf-8")
    (jobs / "project_8.json").write_text(json.dumps({"project_id": 8, "state": "running", "pid": os.getpid()}), encoding="utf-8")
    (jobs / "project_9.json").write_text(json.dumps({"project_id": 9, "state": "done", "pid": 999999999}), encoding="utf-8")

    assert package.fail_stale_package_jobs() == 1
    assert package.read_job(7)["state"] == "error" and "khởi động lại" in package.read_job(7)["message"]
    assert not (jobs / "project_7.lock").exists()
    assert package.read_job(8)["state"] == "running"
    assert package.read_job(9)["state"] == "done"


def test_handover_dir_defaults_next_to_pdf_storage(tmp_path, monkeypatch):
    monkeypatch.delenv("HANDOVER_DIR", raising=False)
    monkeypatch.setenv("PDF_STORAGE_PATH", str(tmp_path / "data" / "uploads"))
    assert package.handover_dir() == (tmp_path / "data" / "handover").resolve()


def test_vb_columns_are_found_by_label():
    columns = package.vb_field_columns(SCHEMA)
    assert columns == {
        "organName": "col_4", "codeNumber": "col_5", "codeNotation": "col_6", "issuedDate": "col_7",
        "typeName": "col_8", "subject": "col_9", "Người ký": "col_10",
    }
    old = [{"category": "x", "fields": [
        {"label": "Số", "name": "col_1"}, {"label": "Số tờ", "name": "col_2"},
        {"label": "Ngày ban hành", "name": "col_3"}, {"label": "Tên loại văn bản", "name": "col_4"},
    ]}]
    assert package.vb_field_columns(old) == {"codeNumber": "col_1", "issuedDate": "col_3", "typeName": "col_4"}


def test_endpoints_are_admin_only_and_report_status(world, monkeypatch):
    for path, method in (("/{project_id}/handover-package", "POST"), ("/{project_id}/handover-package", "GET")):
        route = next(r for r in projects.router.routes if r.path.endswith(path) and method in r.methods)
        assert projects.get_admin_user in {d.call for d in route.dependant.dependencies}
    project_id = world["project"].id
    assert projects.api_get_handover_package(project_id, current_user=ADMIN)["data"]["state"] == "none"

    started = []
    monkeypatch.setattr(package.threading.Thread, "start", lambda self: started.append(self.name))
    response = projects.api_start_handover_package(project_id, current_user=ADMIN, db=world["db"])
    assert response["data"]["state"] == "queued" and started == [f"handover-{project_id}"]
    assert projects.api_get_handover_package(project_id, current_user=ADMIN)["data"]["state"] == "queued"


def report_text(path):
    import re
    import zipfile
    from xml.etree import ElementTree

    with zipfile.ZipFile(path) as archive:
        assert {"[Content_Types].xml", "_rels/.rels", "word/document.xml"} <= set(archive.namelist())
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    paragraphs = []
    for node in root.iter(f"{namespace}p"):
        paragraphs.append("".join(item.text or "" for item in node.iter(f"{namespace}t")))
    return "\n".join(paragraphs)


def test_handover_report_is_written_with_package_figures(world):
    run(world)
    job = package.read_job(world["project"].id)
    path = world["handover"] / ROOT / package.REPORT_FILENAME

    assert job["report_file"] == package.REPORT_FILENAME and path.is_file()
    text = report_text(path)
    sums = world["handover"] / ROOT / "SHA256SUMS.txt"
    for expected in (
        "BIÊN BẢN BÀN GIAO TÀI LIỆU SỐ HÓA", "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", "Dự án: Bộ Y tế 2026",
        "Mã cơ quan: H05.02.02", ROOT, hashlib.sha256(sums.read_bytes()).hexdigest(),
        hashlib.sha256((world["handover"] / ROOT / "Metadata_NN-SIP.xlsx").read_bytes()).hexdigest(),
        "H05.02.02.2006.12", "Tập quyết định 12", "H05.02.02.2006.13a.HC", "ĐẠI DIỆN BÊN GIAO", "ĐẠI DIỆN BÊN NHẬN",
    ):
        assert expected in text, expected
    lines = text.split("\n")  # bảng 2 cột: giá trị ở đoạn ngay sau nhãn
    value_after = lambda label: lines[lines.index(label) + 1]
    assert (value_after("Số hồ sơ"), value_after("Số văn bản"), value_after("Tổng số trang (kể cả bìa)")) == ("2", "3", "10")
    assert "Còn 1 hồ sơ chưa đóng gói" in text


def test_report_counts_failed_dossiers_as_not_packaged(world):
    world["files"]["0001/013a/1.pdf"]["storage"].unlink()
    run(world)
    text = report_text(world["handover"] / ROOT / package.REPORT_FILENAME)
    lines = text.split("\n")

    assert lines[lines.index("Số hồ sơ") + 1] == "1"
    assert "Còn 2 hồ sơ chưa đóng gói" in text
    assert "H05.02.02.2006.13a.HC" not in text


def test_report_download_endpoint(world):
    project_id = world["project"].id
    route = next(r for r in projects.router.routes if r.path.endswith("/handover-package/report"))
    assert projects.get_admin_user in {d.call for d in route.dependant.dependencies}
    with pytest.raises(HTTPException) as error:
        projects.api_download_handover_report(project_id, current_user=ADMIN)
    assert error.value.status_code == 404

    run(world)
    response = projects.api_download_handover_report(project_id, current_user=ADMIN)
    assert Path(response.path) == world["handover"] / ROOT / package.REPORT_FILENAME
    assert response.media_type.endswith("wordprocessingml.document")
    assert f"Bien_ban_ban_giao_du_an_{project_id}.docx" in response.headers["content-disposition"]

    (world["handover"] / ROOT / package.REPORT_FILENAME).unlink()
    with pytest.raises(HTTPException) as error:
        projects.api_download_handover_report(project_id, current_user=ADMIN)
    assert error.value.status_code == 404 and "đóng gói lại" in error.value.detail


def test_docx_writer_escapes_text_and_keeps_line_breaks():
    from server.services.docx_writer import build_docx, paragraph, table

    content = build_docx([paragraph("A <b> & \"c\"\nhai"), table(["x"], [["<y>"]], [1000])])
    import zipfile
    from xml.etree import ElementTree

    with zipfile.ZipFile(BytesIO(content)) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    ElementTree.fromstring(xml)
    assert "A &lt;b&gt; &amp; \"c\"" in xml and "<w:br/>" in xml and "&lt;y&gt;" in xml
