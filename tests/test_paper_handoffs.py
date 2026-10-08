import pytest
import datetime
from fastapi.testclient import TestClient
from sqlalchemy import text
from server.main import app
from server.database import get_db, SessionLocal
from server.models import User, Project, ProjectCase, ProjectStage, ProjectStageMember, Template
from server.models_paper import CasePaperHandoff
from server.migration_runner import HEAD_REVISION, upgrade_database
from openpyxl import load_workbook
import io

def seed_project_and_cases(db, user):
    t1 = Template(name="T1", filename="t1")
    db.add(t1)
    db.commit()
    project = Project(
        name="Test Proj",
        root_folder_name="P1", template_id=t1.id, 
        template_name_snapshot="T1", template_filename_snapshot="t1", 
        case_level=1, report_mode="pdf",
        created_by_user_id=user.id,
        status="ongoing",
    )
    db.add(project)
    db.commit()

    case1 = ProjectCase(project_id=project.id, case_key="box-1", display_name="Hộp 1")
    case2 = ProjectCase(project_id=project.id, case_key="box-2", display_name="Hộp 2")
    db.add_all([case1, case2])
    db.commit()

    return project, case1, case2

@pytest.fixture(scope="function")
def engine_with_fks(tmp_path):
    from sqlalchemy import create_engine
    from sqlalchemy import event
    db_path = tmp_path / "test_fks.db"
    engine = create_engine(f"sqlite:///{db_path}")
    event.listen(engine, "connect", lambda conn, _: conn.execute("PRAGMA foreign_keys=ON"))
    return engine

@pytest.fixture(scope="function")
def test_db(engine_with_fks):
    from alembic.config import Config
    from alembic import command
    alembic_cfg = Config("alembic.ini")
    with engine_with_fks.begin() as connection:
        alembic_cfg.attributes["connection"] = connection
        command.upgrade(alembic_cfg, "head")

    Session = type(SessionLocal)(bind=engine_with_fks)
    db = Session()
    try:
        yield db
    finally:
        db.rollback()
        db.close()
        engine_with_fks.dispose()

@pytest.fixture
def auth_client(test_db):
    client = TestClient(app)
    user = User(username="admin", password="x", full_name="Ad Min", role="admin", account_type="staff")
    test_db.add(user)
    test_db.commit()
    user_id = user.id
    
    app.dependency_overrides[get_db] = lambda: test_db
    from server.routers.auth import get_current_user
    app.dependency_overrides[get_current_user] = lambda: {"id": user_id, "username": "admin", "full_name": "Ad Min", "role": "admin", "session_id": "xyz"}
    yield client
    app.dependency_overrides.clear()

def test_1_record_milestone_and_get(auth_client, test_db):
    user = test_db.query(User).first()
    project, case1, case2 = seed_project_and_cases(test_db, user)

    time_str = "2026-10-08T10:00:00+07:00"
    data = {
        "happened_at": time_str,
        "handed_by": "Nguyễn A",
        "received_by": "Trần B",
        "note": "Ghi chú 1"
    }

    res = auth_client.put(f"/api/projects/{project.id}/cases/{case1.id}/paper-handoffs/received_from_client", json=data)
    assert res.status_code == 200
    assert "Z" in res.json()["data"]["happened_at"]

    res_get = auth_client.get(f"/api/projects/{project.id}/paper-handoffs")
    assert res_get.status_code == 200
    cases = res_get.json()["data"]["cases"]
    assert len(cases) == 2
    case_data = next((c for c in cases if c["case_id"] == case1.id), None)
    assert case_data is not None
    events = case_data["events"]
    assert "received_from_client" in events
    event = events["received_from_client"]
    assert "Z" in event["happened_at"]
    assert event["handed_by"] == "Nguyễn A"

def test_2_missing_previous_milestone(auth_client, test_db):
    user = test_db.query(User).first()
    project, case1, case2 = seed_project_and_cases(test_db, user)

    time_str = "2026-10-08T10:00:00+07:00"
    data = {
        "happened_at": time_str,
        "handed_by": "A",
        "received_by": "B"
    }

    res = auth_client.put(f"/api/projects/{project.id}/cases/{case1.id}/paper-handoffs/to_scan", json=data)
    assert res.status_code == 409
    assert res.json()["detail"]["code"] == "previous_milestone_missing"

def test_3_timeline_validation(auth_client, test_db):
    user = test_db.query(User).first()
    project, case1, case2 = seed_project_and_cases(test_db, user)

    # Record 1
    t1 = "2026-10-08T10:00:00+07:00"
    auth_client.put(f"/api/projects/{project.id}/cases/{case1.id}/paper-handoffs/received_from_client", json={
        "happened_at": t1, "handed_by": "A", "received_by": "B"
    })

    # Record 2 earlier than 1 -> milestone_before_previous
    t2_early = "2026-10-08T09:00:00+07:00"
    res = auth_client.put(f"/api/projects/{project.id}/cases/{case1.id}/paper-handoffs/to_arrangement", json={
        "happened_at": t2_early, "handed_by": "A", "received_by": "B"
    })
    assert res.status_code == 409
    assert res.json()["detail"]["code"] == "milestone_before_previous"

    # Record 2 properly
    t2_proper = "2026-10-08T11:00:00+07:00"
    auth_client.put(f"/api/projects/{project.id}/cases/{case1.id}/paper-handoffs/to_arrangement", json={
        "happened_at": t2_proper, "handed_by": "A", "received_by": "B"
    })

    # Update 1 later than 2 -> milestone_after_next
    t1_late = "2026-10-08T12:00:00+07:00"
    res2 = auth_client.put(f"/api/projects/{project.id}/cases/{case1.id}/paper-handoffs/received_from_client", json={
        "happened_at": t1_late, "handed_by": "A", "received_by": "B"
    })
    assert res2.status_code == 409
    assert res2.json()["detail"]["code"] == "milestone_after_next"

def test_4_future_time_and_missing_tz(auth_client, test_db):
    user = test_db.query(User).first()
    project, case1, case2 = seed_project_and_cases(test_db, user)

    # Future time
    future_time = (datetime.datetime.now(datetime.timezone.utc) + datetime.timedelta(minutes=10)).isoformat()
    res = auth_client.put(f"/api/projects/{project.id}/cases/{case1.id}/paper-handoffs/received_from_client", json={
        "happened_at": future_time, "handed_by": "A", "received_by": "B"
    })
    assert res.status_code == 400
    assert "future_time" in str(res.json())

    # Missing tz
    res2 = auth_client.put(f"/api/projects/{project.id}/cases/{case1.id}/paper-handoffs/received_from_client", json={
        "happened_at": "2026-10-08T10:00:00", "handed_by": "A", "received_by": "B"
    })
    assert res2.status_code == 400

    # Missing person
    res3 = auth_client.put(f"/api/projects/{project.id}/cases/{case1.id}/paper-handoffs/received_from_client", json={
        "happened_at": "2026-10-08T10:00:00+07:00", "handed_by": "", "received_by": "B"
    })
    assert res3.status_code in [400, 422]

def test_5_permissions(test_db):
    admin = User(username="admin5", password="x", role="admin", account_type="staff")
    ctv = User(username="ctv5", password="x", role="user", account_type="ctv")
    staff_scan = User(username="staff_scan", password="x", role="user", account_type="staff")
    staff_other = User(username="staff_other", password="x", role="user", account_type="staff")
    test_db.add_all([admin, ctv, staff_scan, staff_other])
    test_db.commit()

    t1 = Template(name="T1", filename="t1")
    test_db.add(t1)
    test_db.commit()
    project = Project(name="Proj 5", root_folder_name="P1", template_id=t1.id, template_name_snapshot="T1", template_filename_snapshot="t1", case_level=1, report_mode="pdf", created_by_user_id=admin.id, status="ongoing")
    test_db.add(project)
    test_db.commit()

    case = ProjectCase(project_id=project.id, case_key="box-5", display_name="Hộp 5")
    stage_scan = ProjectStage(project_id=project.id, stage_key="scan", position=1)
    stage_data = ProjectStage(project_id=project.id, stage_key="data_entry", position=2)
    test_db.add_all([case, stage_scan, stage_data])
    test_db.commit()

    sm_ctv = ProjectStageMember(project_id=project.id, user_id=ctv.id, stage_key="scan")
    sm_staff = ProjectStageMember(project_id=project.id, user_id=staff_scan.id, stage_key="scan")
    sm_other = ProjectStageMember(project_id=project.id, user_id=staff_other.id, stage_key="data_entry")
    test_db.add_all([sm_ctv, sm_staff, sm_other])
    test_db.commit()

    def do_req(user_obj):
        client = TestClient(app)
        app.dependency_overrides[get_db] = lambda: test_db
        from server.routers.auth import get_current_user
        app.dependency_overrides[get_current_user] = lambda: {"id": user_obj.id, "username": user_obj.username, "role": user_obj.role}
        res = client.put(f"/api/projects/{project.id}/cases/{case.id}/paper-handoffs/received_from_client", json={
            "happened_at": "2026-10-08T10:00:00+07:00", "handed_by": "A", "received_by": "B"
        })
        app.dependency_overrides.clear()
        return res.status_code

    assert do_req(ctv) == 403
    assert do_req(staff_other) == 403
    assert do_req(staff_scan) == 200
    assert do_req(admin) == 200

def test_6_not_found(auth_client):
    res = auth_client.put(f"/api/projects/9999/cases/9999/paper-handoffs/received_from_client", json={
        "happened_at": "2026-10-08T10:00:00+07:00", "handed_by": "A", "received_by": "B"
    })
    assert res.status_code == 404

    res2 = auth_client.put(f"/api/projects/1/cases/1/paper-handoffs/unknown_milestone", json={
        "happened_at": "2026-10-08T10:00:00+07:00", "handed_by": "A", "received_by": "B"
    })
    assert res2.status_code == 404

def test_7_delete(auth_client, test_db):
    user = test_db.query(User).first()
    project, case1, case2 = seed_project_and_cases(test_db, user)
    
    staff = User(username="staff7", password="x", role="user", account_type="staff")
    test_db.add(staff)
    test_db.commit()
    sm = ProjectStageMember(project_id=project.id, user_id=staff.id, stage_key="scan")
    test_db.add(sm)
    test_db.commit()

    auth_client.put(f"/api/projects/{project.id}/cases/{case1.id}/paper-handoffs/received_from_client", json={
        "happened_at": "2026-10-08T10:00:00+07:00", "handed_by": "A", "received_by": "B"
    })
    auth_client.put(f"/api/projects/{project.id}/cases/{case1.id}/paper-handoffs/to_arrangement", json={
        "happened_at": "2026-10-08T11:00:00+07:00", "handed_by": "A", "received_by": "B"
    })

    # Staff delete -> 403
    staff_client = TestClient(app)
    from server.routers.auth import get_current_user
    app.dependency_overrides[get_current_user] = lambda: {"id": staff.id, "username": "staff7", "role": "user"}
    res = staff_client.delete(f"/api/projects/{project.id}/cases/{case1.id}/paper-handoffs/to_arrangement")
    assert res.status_code == 403
    
    # Restore admin override
    app.dependency_overrides[get_current_user] = lambda: {"id": user.id, "username": user.username, "full_name": user.full_name, "role": user.role, "session_id": "xyz"}

    # Delete middle -> 409
    res2 = auth_client.delete(f"/api/projects/{project.id}/cases/{case1.id}/paper-handoffs/received_from_client")
    assert res2.status_code == 409
    assert res2.json()["detail"]["code"] == "not_last_milestone"

    # Delete last -> ok
    res3 = auth_client.delete(f"/api/projects/{project.id}/cases/{case1.id}/paper-handoffs/to_arrangement")
    assert res3.status_code == 200

def test_8_excel(auth_client, test_db):
    user = test_db.query(User).first()
    project, case1, case2 = seed_project_and_cases(test_db, user)

    auth_client.put(f"/api/projects/{project.id}/cases/{case1.id}/paper-handoffs/received_from_client", json={
        "happened_at": "2026-10-08T08:30:00+07:00", "handed_by": "Giao 1", "received_by": "Nhận 1"
    })

    res = auth_client.get(f"/api/projects/{project.id}/paper-handoffs.xlsx")
    assert res.status_code == 200
    
    wb = load_workbook(filename=io.BytesIO(res.content))
    ws = wb["Sổ giao nhận"]
    
    # Header
    headers = [cell.value for cell in ws[1]]
    assert "Hộp" in headers
    assert "Nhận từ khách hàng - Thời gian" in headers
    
    # Rows
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 3 # 1 header, 2 cases
    
    case1_row = [r for r in rows if r[0] == "1"][0]
    assert "08/10/2026 08:30" in case1_row

def test_9_pragma_foreign_keys(test_db):
    user = User(username="admin9", password="x", role="admin", account_type="staff")
    test_db.add(user)
    test_db.commit()
    t1 = Template(name="T1", filename="t1")
    test_db.add(t1)
    test_db.commit()
    project = Project(name="Proj 9", root_folder_name="P1", template_id=t1.id, template_name_snapshot="T1", template_filename_snapshot="t1", case_level=1, report_mode="pdf", created_by_user_id=user.id, status="ongoing")
    test_db.add(project)
    test_db.commit()
    case1 = ProjectCase(project_id=project.id, case_key="box-1", display_name="Hộp 1")
    test_db.add(case1)
    test_db.commit()

    handoff = CasePaperHandoff(project_id=project.id, case_id=case1.id, milestone="received_from_client", 
                               happened_at=datetime.datetime.now(), handed_by="A", received_by="B", recorded_by_user_id=user.id)
    test_db.add(handoff)
    test_db.commit()
    
    test_db.delete(case1)
    test_db.commit()
    
    count = test_db.query(CasePaperHandoff).count()
    assert count == 0

def test_10_migration_additive(engine_with_fks):
    assert HEAD_REVISION == "0016_case_paper_handoffs"
    from alembic.config import Config
    from alembic import command
    alembic_cfg = Config("alembic.ini")
    with engine_with_fks.begin() as connection:
        alembic_cfg.attributes["connection"] = connection
        command.upgrade(alembic_cfg, "head")

    Session = type(SessionLocal)(bind=engine_with_fks)
    db = Session()
    res = db.execute(text("SELECT count(*) FROM case_paper_handoffs")).scalar()
    assert res == 0
    db.close()
