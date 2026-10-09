"""Chi trả theo sản lượng (lát P1a): đơn giá và bảng tạm tính (không lưu)."""
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, inspect
from sqlalchemy.orm import sessionmaker

from server.database import Base, get_db
from server.main import app
from server.models import (
    ArrangementDossier,
    AssignedDocument,
    Project,
    ProjectCase,
    ProjectDocumentAsset,
    ProjectPolicy,
    ProjectReportUnit,
    Submission,
    SubmissionQualityAssessment,
    SubmissionReviewHistory,
    Template,
    User,
)
from server.models_payroll import ProjectWorkRate
from server.models_scan import CaseScanPackage
from server.routers.auth import get_current_user
from server.services.payroll_service import compute_payroll

PERIOD = (date(2026, 10, 1), date(2026, 10, 7))


@pytest.fixture()
def db(tmp_path):
    engine = create_engine(f"sqlite:///{(tmp_path / 'payroll.sqlite3').as_posix()}")
    Base.metadata.create_all(bind=engine)
    session = sessionmaker(bind=engine)()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


@pytest.fixture()
def world(db):
    users = {
        "admin": User(username="admin", password="x", role="admin", full_name="Quản trị", account_type="staff"),
        "nhap": User(username="nhap", password="x", role="user", full_name="Nguyễn Nhập", account_type="staff"),
        "ctv": User(username="ctv", password="x", role="user", full_name="An CTV", account_type="ctv"),
        "duyet": User(username="duyet", password="x", role="user", full_name="Bùi Duyệt", account_type="staff"),
        "scan": User(username="scan", password="x", role="user", full_name="Trần Scan", account_type="staff"),
    }
    template = Template(name="T", filename="t.xlsx")
    db.add_all([*users.values(), template])
    db.flush()
    project = Project(
        name="Bộ Y tế", root_folder_name="Goc", template_id=template.id, template_name_snapshot="T",
        template_filename_snapshot="t.xlsx", case_level=1, report_mode="pdf", created_by_user_id=users["admin"].id, status="ongoing",
    )
    db.add(project)
    db.flush()
    c1 = ProjectCase(project_id=project.id, case_key="0001", display_name="0001")
    c2 = ProjectCase(project_id=project.id, case_key="0002", display_name="0002")
    db.add_all([c1, c2])
    db.flush()

    def dossier(case_row, number, bad=False):
        db.add(ArrangementDossier(
            project_id=project.id, case_id=case_row.id, box_number=int(case_row.case_key), dossier_number=number,
            dossier_suffix="", fonds_code="P", fonds_name="Phông", catalog_number="01", title=f"HS {number}",
            start_date="01/01/2020", end_date="31/12/2020", start_year=2020, maintenance_code="01", sheet_count=10,
            source_row=number, bad_paper=bad,
        ))

    dossier(c1, 1); dossier(c1, 2, bad=True); dossier(c2, 1)
    units = {}
    for case_row in (c1, c2):
        units[case_row.id] = ProjectReportUnit(project_id=project.id, case_id=case_row.id, report_key=case_row.case_key, display_name=case_row.case_key)
        db.add(units[case_row.id])
    db.flush()

    def document(case_row, path, status, *, by=None, submitted_at=None, reviews=()):
        doc = AssignedDocument(original_filename=path, uuid_filename=f"u-{path}")
        db.add(doc)
        db.flush()
        db.add(ProjectDocumentAsset(
            project_id=project.id, case_id=case_row.id, report_unit_id=units[case_row.id].id, assigned_document_id=doc.id,
            relative_path=path, normalized_relative_path=path.casefold(), original_filename=path.rsplit("/", 1)[-1],
            storage_filename=f"s-{path.replace('/', '-')}", content_sha256="0" * 64, byte_size=1,
        ))
        submission = Submission(template_id=template.id, assigned_document_id=doc.id, data_json="{}", status=status,
                                created_by_user_id=users["nhap"].id, created_at=datetime(2026, 9, 1))
        db.add(submission)
        db.flush()
        if submitted_at:
            db.add(SubmissionQualityAssessment(submission_id=submission.id, input_user_id=by.id, baseline_data_json="{}", created_at=submitted_at))
        for reviewed_at in reviews:
            db.add(SubmissionReviewHistory(
                submission_id=submission.id, event_type="review_confirmed", input_user_id=by.id,
                reviewer_user_id=users["duyet"].id, baseline_data_json="{}", reviewer_data_json="{}", created_at=reviewed_at,
            ))
            db.flush()

    # 17:00 UTC ngày 30/09 = 00:00 ngày 01/10 giờ Việt Nam: tính
    document(c1, "0001/001/1.pdf", "completed", by=users["nhap"], submitted_at=datetime(2026, 9, 30, 17, 0),
             reviews=(datetime(2026, 10, 3, 3, 0), datetime(2026, 10, 9, 0, 0)))  # lần duyệt thứ 2 không tính
    # hồ sơ 2 giấy xấu; duyệt 16:59 UTC ngày 07/10 = 23:59 giờ VN: tính
    document(c1, "0001/002/1.pdf", "completed", by=users["ctv"], submitted_at=datetime(2026, 10, 2, 2, 0),
             reviews=(datetime(2026, 10, 7, 16, 59),))
    # nộp 17:00 UTC ngày 07/10 = 00:00 ngày 08/10 giờ VN: ngoài kỳ
    document(c1, "0001/002/2.pdf", "pending_review", by=users["nhap"], submitted_at=datetime(2026, 10, 7, 17, 0))
    document(c1, "0001/001/BIA.pdf", "completed", by=users["nhap"], submitted_at=datetime(2026, 10, 2), reviews=(datetime(2026, 10, 3),))
    document(c2, "0002/001/1.pdf", "draft", by=users["nhap"], submitted_at=datetime(2026, 10, 2))
    # 16:59 UTC ngày 30/09 = 23:59 ngày 30/09 giờ VN: ngoài kỳ
    document(c2, "0002/001/2.pdf", "completed", by=users["nhap"], submitted_at=datetime(2026, 9, 30, 16, 59))

    def package(case_row, version, status, a4, finished_at, scanner=users["scan"]):
        db.add(CaseScanPackage(case_id=case_row.id, version=version, submitted_by_user_id=users["admin"].id,
                               scanned_by_user_id=scanner.id if scanner else None, source_path="x",
                               total_a4_equivalent=a4, status=status, finished_at=finished_at))

    package(c1, 1, "done", 1000, datetime(2026, 10, 2, 3, 0))  # hộp 1 có hồ sơ giấy xấu -> SC-A4-2
    package(c1, 2, "done", 777, datetime(2026, 10, 7, 17, 0))  # 00:00 ngày 08/10 giờ VN: ngoài kỳ
    package(c2, 1, "done", 500, datetime(2026, 10, 5, 3, 0))
    package(c2, 2, "failed", 999, datetime(2026, 10, 6, 3, 0))
    package(c2, 3, "done", 300, datetime(2026, 10, 6, 4, 0), scanner=None)
    for code, price in (("NL-1", 1000), ("CN-1", 400), ("SC-A4-1", 50)):
        db.add(ProjectWorkRate(project_id=project.id, work_code=code, unit_price=Decimal(price)))
    db.commit()
    return project, users


def as_user(user):
    return {"id": user.id, "username": user.username, "full_name": user.full_name, "role": user.role, "session_id": "s"}


def summary(result):
    return [(line["name"], line["work_code"], line["quantity"], line["unit_price"], line["factor"], line["amount"]) for line in result["lines"]]


def test_each_work_code_and_bad_paper_factor(db, world):
    project, _users = world
    result = compute_payroll(db, project_id=project.id, date_from=PERIOD[0], date_to=PERIOD[1])
    assert summary(result) == [
        ("An CTV", "NL-2", 1, 1000, 1.3, 1300),
        ("Bùi Duyệt", "CN-1", 1, 400, 1, 400),
        ("Bùi Duyệt", "CN-2", 1, 400, 1.3, 520),
        ("Nguyễn Nhập", "NL-1", 1, 1000, 1, 1000),
        ("Trần Scan", "SC-A4-1", 500, 50, 1, 25000),
        ("Trần Scan", "SC-A4-2", 1000, 50, 1.3, 65000),
        ("Chưa xác định người scan", "SC-A4-1", 300, 50, 1, 15000),
    ]
    assert result["total"] == 108220
    assert result["warnings"] == []
    assert result["params"] == {"rates": {"NL-1": 1000, "CN-1": 400, "SC-A4-1": 50}, "bad_paper_factor": 1.3, "qc_version": result["params"]["qc_version"]}
    assert [(p["name"], p["amount"], p["complete"]) for p in result["people"]][:2] == [("An CTV", 1300, True), ("Bùi Duyệt", 920, True)]


def test_project_factor_changes_type_two_amounts(db, world):
    project, _users = world
    db.add(ProjectPolicy(project_id=project.id, bad_paper_factor=Decimal("2.00")))
    db.commit()
    lines = {(line["name"], line["work_code"]): line for line in compute_payroll(db, project_id=project.id, date_from=PERIOD[0], date_to=PERIOD[1])["lines"]}
    assert (lines[("An CTV", "NL-2")]["factor"], lines[("An CTV", "NL-2")]["amount"]) == (2, 2000)
    assert lines[("Trần Scan", "SC-A4-2")]["amount"] == 100000
    assert lines[("Nguyễn Nhập", "NL-1")]["amount"] == 1000


def test_missing_rate_warns_without_error(db, world):
    project, _users = world
    db.query(ProjectWorkRate).filter(ProjectWorkRate.work_code == "CN-1").delete()
    db.commit()
    result = compute_payroll(db, project_id=project.id, date_from=PERIOD[0], date_to=PERIOD[1])
    assert result["warnings"] == ["Chưa có đơn giá CN-1"]
    cn = [line for line in result["lines"] if line["work_code"].startswith("CN")]
    assert [(line["unit_price"], line["amount"]) for line in cn] == [(None, None), (None, None)]
    duyet = next(person for person in result["people"] if person["name"] == "Bùi Duyệt")
    assert (duyet["amount"], duyet["complete"]) == (0, False)


def test_period_edges_follow_vietnam_days(db, world):
    project, _users = world
    early = compute_payroll(db, project_id=project.id, date_from=date(2026, 9, 30), date_to=date(2026, 9, 30))
    assert summary(early) == [("Nguyễn Nhập", "NL-1", 1, 1000, 1, 1000)], "chỉ văn bản nộp 23:59 ngày 30/09 giờ VN"
    late = compute_payroll(db, project_id=project.id, date_from=date(2026, 10, 8), date_to=date(2026, 10, 9))
    assert sorted((line["work_code"], line["quantity"]) for line in late["lines"]) == [("NL-2", 1), ("SC-A4-2", 777)], \
        "duyệt lại lần 2 (09/10) không tính; nộp và gói S lúc 00:00 ngày 08/10 giờ VN tính vào 08/10"


def test_rates_api_validation_and_permissions(db, world):
    project, users = world
    client = TestClient(app)
    app.dependency_overrides[get_db] = lambda: db
    try:
        app.dependency_overrides[get_current_user] = lambda: as_user(users["admin"])
        url = f"/api/projects/{project.id}/work-rates"
        assert [rate["unit_price"] for rate in client.get(url).json()["data"]["rates"]] == [1000, 400, 50]
        response = client.put(url, json={"rates": {"NL-1": -1}})
        assert response.status_code == 400 and response.json()["detail"]["code"] == "negative_rate"
        assert client.put(url, json={"rates": {"NL-1": "abc"}}).json()["detail"]["code"] == "invalid_rate"
        assert client.put(url, json={"rates": {"XX-1": 5}}).json()["detail"]["code"] == "unknown_work_code"
        data = client.put(url, json={"rates": {"NL-1": "1250.5", "CN-1": None}}).json()["data"]
        assert [rate["unit_price"] for rate in data["rates"]] == [1250.5, None, 50]
        assert data["bad_paper_factor"] == 1.3
        db.expire_all()
        assert db.query(ProjectWorkRate).filter(ProjectWorkRate.work_code == "NL-1").one().updated_by_user_id == users["admin"].id

        preview = f"/api/projects/{project.id}/payroll-preview"
        assert client.get(preview, params={"from": "2026-10-07", "to": "2026-10-01"}).json()["detail"]["code"] == "period_invalid"
        assert client.get(preview, params={"from": "2026-01-01", "to": "2026-04-03"}).json()["detail"]["code"] == "period_too_long"
        assert client.get(preview, params={"from": "2026-01-01", "to": "2026-04-02"}).status_code == 200, "đúng 92 ngày được"
        assert client.get(preview, params={"from": "01/10/2026", "to": "2026-10-07"}).status_code == 400
        result = client.get(preview, params={"from": "2026-10-01", "to": "2026-10-07"}).json()["data"]
        assert result["warnings"] == ["Chưa có đơn giá CN-1"]
        response = client.get(preview, params={"from": "2026-10-01", "to": "2026-10-07", "format": "xlsx"})
        workbook = openpyxl.load_workbook(BytesIO(response.content))
        assert workbook.sheetnames == ["Tổng theo người", "Chi tiết"]
        assert [cell.value for cell in workbook["Chi tiết"][1]] == ["Người", "Mã công việc", "Công việc", "Đơn vị", "Sản lượng", "Đơn giá", "Hệ số", "Thành tiền"]
        assert response.headers["content-disposition"] == f'attachment; filename="Tam_tinh_chi_tra_{project.id}_2026-10-01_2026-10-07.xlsx"'

        app.dependency_overrides[get_current_user] = lambda: as_user(users["nhap"])
        assert client.get(url).status_code == 403
        assert client.put(url, json={"rates": {"NL-1": 1}}).status_code == 403
        assert client.get(preview, params={"from": "2026-10-01", "to": "2026-10-07"}).status_code == 403
    finally:
        app.dependency_overrides.clear()


def test_work_rates_revision_is_additive():
    pytest.importorskip("alembic")
    from pathlib import Path

    from alembic import command

    from server.migration_runner import HEAD_REVISION, _alembic_config, current_database_revision

    assert HEAD_REVISION == "0019_payroll_periods"
    engine = create_engine("sqlite+pysqlite:///:memory:")
    config = _alembic_config(Path(__file__).resolve().parents[1])
    with engine.begin() as connection:
        config.attributes["connection"] = connection
        command.upgrade(config, "0017_project_handover_lock")
        assert "project_work_rates" not in inspect(connection).get_table_names()
        command.upgrade(config, "0018_project_work_rates")
        assert current_database_revision(connection) == "0018_project_work_rates"
        assert {"project_id", "work_code", "unit_price"} <= {c["name"] for c in inspect(connection).get_columns("project_work_rates")}
        command.downgrade(config, "0017_project_handover_lock")
        assert "project_work_rates" not in inspect(connection).get_table_names()
        command.upgrade(config, "0018_project_work_rates")
        assert "project_work_rates" in inspect(connection).get_table_names()


# ---------------------------------------------------------------- P2: chốt kỳ
def admin_client(db, users):
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: as_user(users["admin"])
    return TestClient(app)


def detail_rows(content):
    sheet = openpyxl.load_workbook(BytesIO(content))["Chi tiết"]
    return [list(row) for row in sheet.iter_rows(min_row=2, values_only=True)]


def test_closed_period_keeps_stored_numbers_after_rates_change(db, world):
    project, users = world
    client = admin_client(db, users)
    try:
        url = f"/api/projects/{project.id}/payroll-periods"
        response = client.post(url, json={"from": "2026-10-01", "to": "2026-10-07"})
        assert response.status_code == 200, response.text
        period = response.json()["data"]
        assert (period["total"], period["lines"]) == (108220, 7)
        assert period["params"]["rates"] == {"NL-1": 1000, "CN-1": 400, "SC-A4-1": 50}
        assert period["params"]["bad_paper_factor"] == 1.3 and period["params"]["qc_version"]
        before = client.get(f"{url}/{period['id']}.xlsx")
        assert before.headers["content-disposition"] == f'attachment; filename="Chi_tra_{project.id}_2026-10-01_2026-10-07.xlsx"'

        client.put(f"/api/projects/{project.id}/work-rates", json={"rates": {"NL-1": 9999, "SC-A4-1": 1}})
        db.add(ProjectPolicy(project_id=project.id, bad_paper_factor=Decimal("3.00")))
        db.commit()
        after = client.get(f"{url}/{period['id']}.xlsx")
        assert detail_rows(after.content) == detail_rows(before.content), "kỳ đã chốt không đổi theo đơn giá mới"
        assert ["An CTV", "NL-2", "Nhập liệu", "văn bản", 1, 1000, 1.3, 1300] in detail_rows(after.content)
        summary = [list(row)[:2] for row in openpyxl.load_workbook(BytesIO(after.content))["Tổng theo người"].iter_rows(values_only=True)]
        assert ["Đơn giá loại 1", "NL-1: 1000, CN-1: 400, SC-A4-1: 50"] in summary, "tham số lúc chốt"
        assert ["Hệ số giấy xấu", 1.3] in summary
        listed = client.get(url).json()["data"]
        assert [(item["id"], item["from"], item["to"], item["total"], item["created_by"]) for item in listed] == [
            (period["id"], "2026-10-01", "2026-10-07", 108220, "Quản trị")]
        assert client.get(f"{url}/999.xlsx").status_code == 404
    finally:
        app.dependency_overrides.clear()


def test_overlap_and_missing_rate_block_closing(db, world):
    project, users = world
    client = admin_client(db, users)
    try:
        url = f"/api/projects/{project.id}/payroll-periods"
        assert client.post(url, json={"from": "2026-10-01", "to": "2026-10-07"}).status_code == 200
        for start, end in (("2026-10-07", "2026-10-10"), ("2026-09-25", "2026-10-01"), ("2026-10-03", "2026-10-04")):
            response = client.post(url, json={"from": start, "to": end})
            assert response.status_code == 409 and response.json()["detail"]["code"] == "period_overlap", (start, end)
        assert client.post(url, json={"from": "2026-10-08", "to": "2026-10-10"}).status_code == 200, "kỳ liền kề không chồng"
        db.query(ProjectWorkRate).filter(ProjectWorkRate.work_code == "SC-A4-1").delete()
        db.commit()
        response = client.post(url, json={"from": "2026-09-01", "to": "2026-09-30"})
        assert response.status_code == 200, "kỳ không có sản lượng scan thì không cần đơn giá scan"
        response = client.post(url, json={"from": "2026-11-01", "to": "2026-11-02"})
        assert response.status_code == 200
        db.query(ProjectWorkRate).delete()
        db.commit()
        db.add(CaseScanPackage(case_id=db.query(ProjectCase).first().id, version=9, submitted_by_user_id=users["admin"].id,
                               source_path="x", total_a4_equivalent=5, status="done", finished_at=datetime(2026, 12, 2, 3)))
        db.commit()
        response = client.post(url, json={"from": "2026-12-01", "to": "2026-12-03"})
        assert response.status_code == 409 and response.json()["detail"]["code"] == "missing_rate"
        assert "Chưa có đơn giá SC-A4-1" in response.json()["detail"]["message"]
        assert len(client.get(url).json()["data"]) == 4, "thiếu đơn giá thì không lưu kỳ"
    finally:
        app.dependency_overrides.clear()


def test_only_latest_period_can_be_deleted(db, world):
    project, users = world
    client = admin_client(db, users)
    try:
        url = f"/api/projects/{project.id}/payroll-periods"
        first = client.post(url, json={"from": "2026-10-01", "to": "2026-10-07"}).json()["data"]["id"]
        second = client.post(url, json={"from": "2026-10-08", "to": "2026-10-09"}).json()["data"]["id"]
        response = client.delete(f"{url}/{first}")
        assert response.status_code == 409 and response.json()["detail"]["code"] == "not_latest_period"
        assert client.delete(f"{url}/{second}").status_code == 200
        from server.models_payroll import PayrollLine, PayrollPeriod
        db.expire_all()
        assert db.query(PayrollPeriod).count() == 1
        assert db.query(PayrollLine).filter(PayrollLine.period_id == second).count() == 0, "xóa kỳ xóa cả dòng"
        assert client.delete(f"{url}/{first}").status_code == 200
        app.dependency_overrides[get_current_user] = lambda: as_user(users["nhap"])
        assert client.post(url, json={"from": "2026-10-01", "to": "2026-10-07"}).status_code == 403
        assert client.get(url).status_code == 403
    finally:
        app.dependency_overrides.clear()


def test_deleting_period_cascades_with_foreign_keys_on(tmp_path):
    from sqlalchemy import event, text

    from server.models_payroll import PayrollLine, PayrollPeriod

    engine = create_engine(f"sqlite:///{(tmp_path / 'fk.sqlite3').as_posix()}")
    event.listen(engine, "connect", lambda connection, _record: connection.execute("PRAGMA foreign_keys=ON"))
    Base.metadata.create_all(bind=engine)
    session = sessionmaker(bind=engine)()
    try:
        user = User(username="a", password="x", role="admin", full_name="A")
        template = Template(name="T", filename="t.xlsx")
        session.add_all([user, template])
        session.flush()
        project = Project(name="P", root_folder_name="P", template_id=template.id, template_name_snapshot="T",
                          template_filename_snapshot="t.xlsx", case_level=1, report_mode="pdf", created_by_user_id=user.id)
        session.add(project)
        session.flush()
        period = PayrollPeriod(project_id=project.id, date_from=date(2026, 10, 1), date_to=date(2026, 10, 2), params_json="{}",
                               total_amount=0, created_by_user_id=user.id)
        session.add(period)
        session.flush()
        session.add(PayrollLine(period_id=period.id, person_name="A", work_code="NL-1", quantity=1, unit_price=1, factor=1, amount=1))
        session.commit()
        project_id, user_id = project.id, user.id
        session.execute(text("DELETE FROM payroll_periods"))
        session.commit()
        session.expunge_all()
        assert session.query(PayrollLine).count() == 0
        session.add(PayrollPeriod(project_id=project_id, date_from=date(2026, 10, 1), date_to=date(2026, 10, 2), params_json="{}",
                                  total_amount=0, created_by_user_id=user_id))
        session.commit()
        session.execute(text("DELETE FROM projects"))
        session.commit()
        assert session.query(PayrollPeriod).count() == 0, "xóa dự án xóa kỳ"
    finally:
        session.close()
        engine.dispose()


def test_payroll_periods_revision_is_additive():
    pytest.importorskip("alembic")
    from pathlib import Path

    from alembic import command

    from server.migration_runner import _alembic_config, current_database_revision

    engine = create_engine("sqlite+pysqlite:///:memory:")
    config = _alembic_config(Path(__file__).resolve().parents[1])
    with engine.begin() as connection:
        config.attributes["connection"] = connection
        command.upgrade(config, "0018_project_work_rates")
        assert not {"payroll_periods", "payroll_lines"} & set(inspect(connection).get_table_names())
        command.upgrade(config, "0019_payroll_periods")
        assert current_database_revision(connection) == "0019_payroll_periods"
        assert {"payroll_periods", "payroll_lines"} <= set(inspect(connection).get_table_names())
        command.downgrade(config, "0018_project_work_rates")
        assert not {"payroll_periods", "payroll_lines"} & set(inspect(connection).get_table_names())
        command.upgrade(config, "0019_payroll_periods")
        assert {"payroll_periods", "payroll_lines"} <= set(inspect(connection).get_table_names())
