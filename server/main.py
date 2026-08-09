import os
import sys
import json
import shutil
import uuid
from fastapi import FastAPI, Depends, BackgroundTasks, UploadFile, File
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

# Determine base directory for bundled data files
if getattr(sys, 'frozen', False):
    # Running as PyInstaller bundle: data files are in _MEIPASS
    BASE_DIR = sys._MEIPASS
    # Working directory (for DB, uploads, scratch) is where the exe lives
    WORK_DIR = os.path.dirname(sys.executable)
else:
    BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    WORK_DIR = BASE_DIR

os.makedirs(os.path.join(WORK_DIR, "uploads"), exist_ok=True)
os.makedirs(os.path.join(WORK_DIR, "scratch"), exist_ok=True)

# Add base directory to sys.path to reuse excel_handler
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)
from excel_handler import get_form_schema, get_ma_xa_mapping
from server.models import SessionLocal, Submission

app = FastAPI(title="Scan To Excel Web App (Enterprise)")

TEMPLATE_PATH = os.path.join(BASE_DIR, "tai_lieu_mau", "Excel_FormMau_v5_04082026.xlsx")

# Dependency
def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

class SubmitRequest(BaseModel):
    data: dict

@app.get("/api/schema")
def api_get_schema():
    """Returns the form schema with dropdown options."""
    try:
        schema = get_form_schema(TEMPLATE_PATH)
        return {"status": "ok", "data": schema}
    except Exception as e:
        return {"status": "error", "message": str(e)}

@app.get("/api/maxa_mapping")
def api_get_maxa_mapping():
    """Returns the mapping for communes."""
    try:
        from excel_handler import get_ma_xa_mapping
        mapping = get_ma_xa_mapping(TEMPLATE_PATH)
        return {"status": "ok", "data": mapping}
    except Exception as e:
        return {"status": "error", "message": str(e)}

@app.get("/api/don-vi-do-mapping")
def api_get_don_vi_do_mapping():
    """Returns the mapping for Don vi do dac."""
    try:
        from excel_handler import get_don_vi_do_mapping
        mapping = get_don_vi_do_mapping(TEMPLATE_PATH)
        return {"status": "ok", "data": mapping}
    except Exception as e:
        return {"status": "error", "message": str(e)}

@app.post("/api/submit")
def api_submit(req: SubmitRequest, db: Session = Depends(get_db)):
    """Saves a submission to the SQLite database."""
    try:
        sub = Submission(
            data_json=json.dumps(req.data, ensure_ascii=False)
        )
        db.add(sub)
        db.commit()
        return {"status": "ok"}
    except Exception as e:
        db.rollback()
        return {"status": "error", "message": str(e)}

@app.get("/api/submissions")
def api_get_submissions(db: Session = Depends(get_db)):
    """Returns a list of all saved submissions for display in the table."""
    try:
        submissions = db.query(Submission).order_by(Submission.created_at.desc(), Submission.id.desc()).all()
        results = []
        for sub in submissions:
            data_dict = json.loads(sub.data_json)
            # Try to extract identifying information. Depending on the schema:
            # We look for "col_8" (Họ tên đại diện), "col_25" (Họ tên vợ chồng)...
            # We'll just grab some basic fields or default to "Hồ sơ #"
            ho_ten = data_dict.get("col_8", "") or data_dict.get("col_25", "") or "(Chưa có tên)"
            so_giay_to = data_dict.get("col_13", "") or data_dict.get("col_30", "") or "(Chưa có CMND)"
            pdf_filename = data_dict.get("_pdf_filename", "")
            
            results.append({
                "id": sub.id,
                "created_at": sub.created_at.strftime("%Y-%m-%d %H:%M:%S"),
                "ho_ten": ho_ten,
                "so_giay_to": so_giay_to,
                "pdf_filename": pdf_filename
            })
        return {"status": "ok", "data": results}
    except Exception as e:
        return {"status": "error", "message": str(e)}

@app.get("/api/submissions/by-pdf")
def api_get_submission_by_pdf(filename: str, db: Session = Depends(get_db)):
    """Returns the first submission associated with a given PDF filename."""
    try:
        # Since _pdf_filename is stored inside data_json, we need to query and filter in python,
        # or use JSON operators if sqlite supports them. For simplicity and robustness with small datasets, 
        # we can just fetch all or use LIKE.
        search_str = f'"{filename}"'
        submissions = db.query(Submission).filter(Submission.data_json.contains(search_str)).order_by(Submission.created_at.desc(), Submission.id.desc()).all()
        for sub in submissions:
            data_dict = json.loads(sub.data_json)
            if data_dict.get("_pdf_filename") == filename:
                return {"status": "ok", "data": data_dict, "id": sub.id}
        return {"status": "ok", "data": None}
    except Exception as e:
        return {"status": "error", "message": str(e)}

@app.get("/api/submissions/{sub_id}")
def api_get_submission(sub_id: int, db: Session = Depends(get_db)):
    """Returns details of a specific submission."""
    try:
        sub = db.query(Submission).filter(Submission.id == sub_id).first()
        if not sub:
            return {"status": "error", "message": "Không tìm thấy hồ sơ."}
        return {"status": "ok", "data": json.loads(sub.data_json)}
    except Exception as e:
        return {"status": "error", "message": str(e)}

@app.put("/api/submissions/{sub_id}")
def api_update_submission(sub_id: int, req: SubmitRequest, db: Session = Depends(get_db)):
    """Updates an existing submission."""
    try:
        sub = db.query(Submission).filter(Submission.id == sub_id).first()
        if not sub:
            return {"status": "error", "message": "Không tìm thấy hồ sơ."}
        
        sub.data_json = json.dumps(req.data, ensure_ascii=False)
        db.commit()
        return {"status": "ok"}
    except Exception as e:
        db.rollback()
        return {"status": "error", "message": str(e)}

@app.delete("/api/submissions/{sub_id}")
def api_delete_submission(sub_id: int, db: Session = Depends(get_db)):
    """Deletes a submission."""
    try:
        sub = db.query(Submission).filter(Submission.id == sub_id).first()
        if not sub:
            return {"status": "error", "message": "Không tìm thấy hồ sơ."}
        
        db.delete(sub)
        db.commit()
        return {"status": "ok"}
    except Exception as e:
        db.rollback()
        return {"status": "error", "message": str(e)}

from datetime import timedelta

@app.post("/api/submissions/{sub_id}/copy")
def api_copy_submission(sub_id: int, db: Session = Depends(get_db)):
    """Duplicates a submission and places it immediately below the original."""
    try:
        sub = db.query(Submission).filter(Submission.id == sub_id).first()
        if not sub:
            return {"status": "error", "message": "Không tìm thấy hồ sơ."}
        
        new_created_at = sub.created_at - timedelta(milliseconds=1)
        
        new_sub = Submission(
            data_json=sub.data_json,
            template_name=sub.template_name,
            created_at=new_created_at
        )
        db.add(new_sub)
        db.commit()
        return {"status": "ok"}
    except Exception as e:
        db.rollback()
        return {"status": "error", "message": str(e)}

@app.get("/api/export")
def api_export(background_tasks: BackgroundTasks, db: Session = Depends(get_db), mode: str = "new"):
    """Exports all submissions to the Excel template and returns the file.
    mode='new': creates a fresh report from template.
    mode='append': appends data to the existing report file if it exists.
    """
    try:
        import openpyxl
        from datetime import datetime
        
        os.makedirs("scratch", exist_ok=True)
        
        # Persistent report file for "append" mode
        persistent_path = os.path.join("scratch", "BaoCao_HienTai.xlsx")
        
        if mode == "append" and os.path.exists(persistent_path):
            # Append to existing report
            wb = openpyxl.load_workbook(persistent_path)
        else:
            # Create fresh from template
            shutil.copy(TEMPLATE_PATH, persistent_path)
            wb = openpyxl.load_workbook(persistent_path)
        
        sht_name = next((s for s in wb.sheetnames if s.strip().lower() == 'data'), None)
        if not sht_name:
            raise Exception("Không tìm thấy sheet 'Data' trong file mẫu.")
        ws = wb[sht_name]
        
        if mode != "append":
            # Clear any existing data rows (row 5 onwards) for a clean slate
            if ws.max_row >= 5:
                ws.delete_rows(5, ws.max_row - 4)
        
        def process_value(val):
            if val and " - " in str(val):
                parts = str(val).split(" - ", 1)
                if len(parts) == 2 and parts[0].strip().isdigit():
                    return parts[0].strip()
            return val

        submissions = db.query(Submission).order_by(Submission.id).all()
        if not submissions:
            raise Exception("Không có dữ liệu hồ sơ nào trong hệ thống để xuất báo cáo.")
            
        for sub in submissions:
            data_dict = json.loads(sub.data_json)
            new_row = [""] * (ws.max_column + 10)
            
            for key, value in data_dict.items():
                if key.startswith('col_'):
                    idx = int(key.split('_')[1])
                    new_row[idx] = process_value(value)
                
            ws.append(new_row)
            
        wb.save(persistent_path)
        wb.close()
        
        # Copy to a timestamped download file
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        download_filename = f"BaoCao_{timestamp}.xlsx"
        download_path = os.path.join("scratch", download_filename)
        shutil.copy(persistent_path, download_path)
        
        def remove_file(path):
            try:
                os.remove(path)
            except:
                pass
        
        background_tasks.add_task(remove_file, download_path)
        
        return FileResponse(download_path, filename=download_filename)
    except Exception as e:
        return {"status": "error", "message": str(e)}

@app.post("/api/upload-pdf")
async def api_upload_pdf(file: UploadFile = File(...)):
    """Uploads a PDF or Image to be viewed side-by-side."""
    try:
        os.makedirs("uploads", exist_ok=True)
        # Generate random filename to prevent clashes
        ext = os.path.splitext(file.filename)[1]
        new_filename = f"{uuid.uuid4().hex}{ext}"
        filepath = os.path.join("uploads", new_filename)
        
        with open(filepath, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
            
        return {"status": "ok", "url": f"/uploads/{new_filename}"}
    except Exception as e:
        return {"status": "error", "message": str(e)}

@app.post("/api/export-append")
async def api_export_append(background_tasks: BackgroundTasks, file: UploadFile = File(...), db: Session = Depends(get_db)):
    """Accepts an existing Excel file, appends current DB submissions to it, and returns the updated file."""
    try:
        import openpyxl
        from datetime import datetime
        
        os.makedirs("scratch", exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        upload_path = os.path.join("scratch", f"upload_{timestamp}.xlsx")
        
        # Save the uploaded file
        with open(upload_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
        
        wb = openpyxl.load_workbook(upload_path)
        sht_name = next((s for s in wb.sheetnames if s.strip().lower() == 'data'), None)
        if not sht_name:
            raise Exception("Không tìm thấy sheet 'Data' trong file đã tải lên.")
        ws = wb[sht_name]
        
        def process_value(val):
            if val and " - " in str(val):
                parts = str(val).split(" - ", 1)
                if len(parts) == 2 and parts[0].strip().isdigit():
                    return parts[0].strip()
            return val

        submissions = db.query(Submission).order_by(Submission.id).all()
        if not submissions:
            raise Exception("Không có dữ liệu hồ sơ nào trong hệ thống để thêm vào báo cáo.")
            
        for sub in submissions:
            data_dict = json.loads(sub.data_json)
            new_row = [""] * (ws.max_column + 10)
            
            for key, value in data_dict.items():
                if key.startswith('col_'):
                    idx = int(key.split('_')[1])
                    new_row[idx] = process_value(value)
                
            ws.append(new_row)
            
        result_filename = f"BaoCao_BoSung_{timestamp}.xlsx"
        result_path = os.path.join("scratch", result_filename)
        wb.save(result_path)
        wb.close()
        
        def remove_file(path):
            try:
                os.remove(path)
            except:
                pass
        
        # Clean up uploaded file
        remove_file(upload_path)
        
        background_tasks.add_task(remove_file, result_path)
        
        return FileResponse(result_path, filename=result_filename)
    except Exception as e:
        return {"status": "error", "message": str(e)}

# Serve uploaded files
app.mount("/uploads", StaticFiles(directory=os.path.join(WORK_DIR, "uploads")), name="uploads")

# Serve frontend static files
app.mount("/", StaticFiles(directory=os.path.join(BASE_DIR, "frontend"), html=True), name="frontend")
