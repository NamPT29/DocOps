import pandas as pd
import os
import json
import logging
import shutil
import openpyxl
from copy import copy


logger = logging.getLogger(__name__)


class ExcelTemplateError(ValueError):
    """The workbook cannot be interpreted as an input form."""


class ExcelExportError(RuntimeError):
    """The workbook or submission data cannot be exported safely."""


def _detect_excel_layout(excel_path):
    """Return the worksheet and zero-based header rows used by a template."""
    workbook = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    try:
        data_sheet = next(
            (
                name
                for name in workbook.sheetnames
                if name.strip().lower() == 'data'
                and workbook[name].sheet_state == 'visible'
            ),
            None,
        )
        sheet_name = data_sheet or next(
            (sheet.title for sheet in workbook.worksheets if sheet.sheet_state == 'visible'),
            workbook.sheetnames[0],
        )
    finally:
        workbook.close()

    preview = pd.read_excel(excel_path, sheet_name=sheet_name, header=None, nrows=50)
    row_counts = preview.notna().sum(axis=1).astype(int).tolist()
    candidates = [
        index
        for index in range(len(row_counts) - 1)
        if row_counts[index] > 0
        and row_counts[index + 1] > row_counts[index]
        and (index == 0 or row_counts[index - 1] <= row_counts[index])
    ]
    if not candidates:
        raise ValueError(f"Không nhận diện được các hàng tiêu đề trong sheet '{sheet_name}'.")

    header_start = max(candidates, key=lambda index: row_counts[index + 1] - row_counts[index])
    header_depth = 4 if data_sheet and header_start == 0 else 3
    header_rows = list(
        range(header_start, min(header_start + header_depth, len(row_counts)))
    )
    if len(header_rows) < 2:
        raise ValueError(f"Sheet '{sheet_name}' không có đủ hàng tiêu đề.")
    return sheet_name, header_rows



def _load_excel_headers(excel_path):
    try:
        sheet_name, header_rows = _detect_excel_layout(excel_path)
        df_head = pd.read_excel(
            excel_path,
            sheet_name=sheet_name,
            header=header_rows,
            nrows=0,
        )
        return sheet_name, df_head
    except Exception as exc:
        raise ExcelTemplateError(
            f"Không thể tự nhận diện biểu mẫu Excel: {exc}"
        ) from exc

def _build_field_schema(col_idx, levels, is_legacy_data_sheet, config, dicts):
    levels = [str(value).strip() for value in levels]
    levels = ["" if value.startswith('Unnamed:') else value for value in levels]
    cat = levels[0] if levels else ""
    field1 = levels[1] if len(levels) > 1 else ""
    field2 = levels[2] if len(levels) > 2 else ""
    field3 = levels[3] if len(levels) > 3 else ""
    
    # Skip 'Số TT' as user requested
    if 'số tt' in cat.lower() or 'số tt' in field1.lower():
        return None, None
    
    if not any(levels):
        return None, None

    col_1 = col_idx + 1
    if is_legacy_data_sheet and ((17 <= col_1 <= 22) or (34 <= col_1 <= 39) or (86 <= col_1 <= 89) or (90 <= col_1 <= 94) or (111 <= col_1 <= 168) or (179 <= col_1 <= 182)):
        label_parts = [p for p in [field2, field3] if p]
    else:
        label_parts = [value for value in levels[1:] if value]
        
    label = " - ".join(label_parts) if label_parts else cat
    
    if not cat:
        cat = "Thông tin chung"

    field_type = "text"
    options = []
    extract_mode = "none"
    
    if config and isinstance(config, dict):
        dd_rules = config.get("dropdown_rules", [])
        for rule in dd_rules:
            if str(rule.get("col")) == str(col_1):
                field_type = "dropdown"
                dict_name = rule.get("dictionary")
                options = dicts.get(dict_name, [])
                extract_mode = rule.get("extract_mode", "none")
                break
            
    field = {
        "col_index": col_idx,
        "label": label,
        "name": f"col_{col_idx}",
        "type": field_type,
        "options": options
    }
    if config and isinstance(config, dict):
        seps = config.get("separators", {})
        if str(col_1) in seps:
            field["separator_above"] = seps[str(col_1)]["title"]
            field["group_end"] = int(seps[str(col_1)]["group_end"]) if seps[str(col_1)].get("group_end") else None
    elif is_legacy_data_sheet:
        if col_1 == 17:
            field["separator_above"] = "Địa chỉ sử dụng"
            field["group_end"] = 22
        elif col_1 == 34:
            field["separator_above"] = "Địa chỉ vợ (chồng)"
            field["group_end"] = 39
        elif col_1 == 86:
            field["separator_above"] = "Diện tích hành lang"
            field["group_end"] = 89
        elif col_1 == 90:
            field["separator_above"] = "Địa chỉ thửa đất"
            field["group_end"] = 94
        elif col_1 == 111:
            field["separator_above"] = "Hạn chế quyền"
            field["group_end"] = 117
        elif col_1 == 118:
            field["separator_above"] = "Nghĩa vụ tài chính"
            field["group_end"] = 123
        elif col_1 == 124:
            field["separator_above"] = "Miễn giảm nghĩa vụ tài chính"
            field["group_end"] = 128
        elif col_1 == 129:
            field["separator_above"] = "Nợ nghĩa vụ tài chính"
            field["group_end"] = 133
        elif col_1 == 134:
            field["separator_above"] = "Nhà ở riêng lẻ"
            field["group_end"] = 141
        elif col_1 == 142:
            field["separator_above"] = "Công trình, hạng mục công trình xây dựng"
            field["group_end"] = 154
        elif col_1 == 155:
            field["separator_above"] = "Công trình ngầm"
            field["group_end"] = 162
        elif col_1 == 163:
            field["separator_above"] = "Rừng trồng"
            field["group_end"] = 165
        elif col_1 == 166:
            field["separator_above"] = "Cây lâu năm"
            field["group_end"] = 168
        elif col_1 == 179:
            field["separator_above"] = "Thông tin lưu kho vật lý hồ sơ"
            field["group_end"] = 182
        
    if field_type == "dropdown":
        field["extract_mode"] = extract_mode

    return cat, field

def get_form_schema(excel_path, dicts=None, config=None):
    """
    Detects the data worksheet and its header rows to construct a hierarchical form schema.
    Uses provided dicts or an empty dictionary.
    Optionally applies dynamic `config` (dict) to override hardcoded behaviors.
    """
    sheet_name, df_head = _load_excel_headers(excel_path)
    if dicts is None:
        dicts = {}

    schema = []
    current_category = ""
    category_fields = []
    
    is_legacy_data_sheet = sheet_name.strip().lower() == 'data'
    
    for col_idx, col_tuple in enumerate(df_head.columns):
        levels = list(col_tuple) if isinstance(col_tuple, tuple) else [col_tuple]
        cat, field = _build_field_schema(col_idx, levels, is_legacy_data_sheet, config, dicts)
        
        if not field:
            continue
            
        if cat != current_category:
            if current_category:
                schema.append({
                    "category": current_category,
                    "fields": category_fields
                })
            current_category = cat
            category_fields = []
            
        category_fields.append(field)
        
    if current_category:
        schema.append({
            "category": current_category,
            "fields": category_fields
        })
        
    return schema

def get_ma_xa_mapping(excel_path):
    """
    Extracts 3-level and 2-level mapping for communes from 'MaXa_3cap_2cap' sheet.
    Returns: { "mapping_3_cap": dict, "mapping_2_cap": dict }
    """
    import re
    try:
        df = pd.read_excel(excel_path, sheet_name='MaXa_3cap_2cap', dtype=str)
        mapping_3_cap = {}
        mapping_2_cap = {}
        
        def normalize_str(s):
            if pd.isna(s): return ""
            s = str(s).lower()
            # Sử dụng chung một regex với JS để đảm bảo đồng nhất 100%
            s = re.sub(r'[^a-z0-9áàảãạăắằẳẵặâấầẩẫậéèẻẽẹêếềểễệíìỉĩịóòỏõọôốồổỗộơớờởỡợúùủũụưứừửữựýỳỷỹỵđ]', '', s)
            return s
            
        def clean_code(c):
            c = str(c).strip()
            if c.endswith('.0'): c = c[:-2]
            return c
            
        for _, row in df.iterrows():
            # 3 cap
            if 'Mã xã 3 cấp (cũ)' in row:
                t3_raw = str(row.get('Tên xã 3 cấp')).strip() if not pd.isna(row.get('Tên xã 3 cấp')) else ''
                h3_raw = str(row.get('Quận/huyện 3 cấp')).strip() if not pd.isna(row.get('Quận/huyện 3 cấp')) else ''
                p3_raw = str(row.get('Tỉnh 3 cấp (cũ)')).strip() if not pd.isna(row.get('Tỉnh 3 cấp (cũ)')) else ''
                
                t3 = normalize_str(t3_raw)
                h3 = normalize_str(h3_raw)
                p3 = normalize_str(p3_raw)
                
                if t3 and h3 and p3:
                    addr_3 = f"{t3}{h3}{p3}"
                    code_3 = row['Mã xã 3 cấp (cũ)']
                    if not pd.isna(code_3):
                        mapping_3_cap[addr_3] = {
                            "code": clean_code(code_3),
                            "formatted": f"{t3_raw}, {h3_raw}, {p3_raw}"
                        }
                
            # 2 cap
            if 'Mã xã 2 cấp (mới)' in row:
                t2_raw = str(row.get('Tên xã 2 cấp')).strip() if not pd.isna(row.get('Tên xã 2 cấp')) else ''
                p2_raw = str(row.get('Tỉnh/TP (2 cấp)')).strip() if not pd.isna(row.get('Tỉnh/TP (2 cấp)')) else ''
                
                t2 = normalize_str(t2_raw)
                p2 = normalize_str(p2_raw)
                
                if t2 and p2:
                    addr_2 = f"{t2}{p2}"
                    code_2 = row['Mã xã 2 cấp (mới)']
                    if not pd.isna(code_2):
                        mapping_2_cap[addr_2] = {
                            "code": clean_code(code_2),
                            "formatted": f"{t2_raw}, {p2_raw}"
                        }
                
        return {
            "mapping_3_cap": mapping_3_cap,
            "mapping_2_cap": mapping_2_cap
        }
    except Exception:
        logger.warning("Không thể đọc mapping mã xã từ %s", excel_path, exc_info=True)
        return {"mapping_3_cap": {}, "mapping_2_cap": {}}

def get_don_vi_do_mapping(excel_path):
    """
    Extracts mapping from QNH_ThongTinDoDac sheet to automatically fill Don vi do dac based on address.
    Returns: { "normalized_phuong_huyen": {"don_vi": "Tên đơn vị đo", "ngay_hoan_thanh": "Ngày hoàn thành"} }
    """
    import re
    try:
        df = pd.read_excel(excel_path, sheet_name='QNH_ThongTinDoDac', dtype=str)
        # Handle dynamic columns in case of encoding issues
        col_phuong = next((c for c in df.columns if 'ph' in c.lower() and 'c' in c.lower()), df.columns[1] if len(df.columns) > 1 else 'Tên phường cũ')
        col_huyen = next((c for c in df.columns if 'huy' in c.lower()), df.columns[2] if len(df.columns) > 2 else 'Huyện')
        col_don_vi = next((c for c in df.columns if 'đơn vị' in c.lower() or 'v' in c.lower()), df.columns[3] if len(df.columns) > 3 else 'Tên đơn vị đo')
        col_ngay = next((c for c in df.columns if 'ng' in c.lower() and 'ho' in c.lower()), df.columns[4] if len(df.columns) > 4 else 'Ngày hoàn thành')

        mapping = {}
        for _, row in df.iterrows():
            phuong = str(row.get(col_phuong, '')).strip()
            huyen = str(row.get(col_huyen, '')).strip()
            don_vi = str(row.get(col_don_vi, '')).strip()
            ngay = str(row.get(col_ngay, '')).strip()
            
            if ngay and ngay != 'nan':
                if ngay.replace('.', '', 1).isdigit():
                    try:
                        ngay = pd.to_datetime(float(ngay), origin='1899-12-30', unit='D').strftime("%d/%m/%Y")
                    except (TypeError, ValueError, OverflowError):
                        logger.debug(
                            "Giá trị ngày đo đạc không hợp lệ: %r",
                            ngay,
                            exc_info=True,
                        )
            else:
                ngay = ""

            if phuong and phuong != 'nan' and don_vi and don_vi != 'nan':
                # create a normalized key: lowercase and alphanumeric only
                s = (phuong + huyen).lower()
                key = re.sub(r'[^a-z0-9áàảãạăắằẳẵặâấầẩẫậéèẻẽẹêếềểễệíìỉĩịóòỏõọôốồổỗộơớờởỡợúùủũụưứừửữựýỳỷỹỵđ]', '', s)
                mapping[key] = {
                    "don_vi": don_vi,
                    "ngay_hoan_thanh": ngay
                }
        return mapping
    except Exception:
        logger.warning(
            "Không thể đọc mapping đơn vị đo từ %s",
            excel_path,
            exc_info=True,
        )
        return {}

def export_submissions_to_excel(template_file_path: str, submissions: list, download_path: str):
    """
    Exports a list of submissions into a provided Excel template.
    Returns the path to the exported file.
    """
    workbook = None
    try:
        sheet_name, header_rows = _detect_excel_layout(template_file_path)
        shutil.copy(template_file_path, download_path)
        keep_vba = os.path.splitext(template_file_path)[1].lower() == '.xlsm'
        workbook = openpyxl.load_workbook(download_path, keep_vba=keep_vba)

        worksheet = workbook[sheet_name]
        data_start_row = max(header_rows) + 2

        # Capture only the cell frames before deleting sample/old data. Exported
        # values should use Excel's default font/fill/alignment/number format while
        # retaining the border layout designed by the template. Header rows above
        # data_start_row remain untouched.
        template_column_count = worksheet.max_column
        template_borders = [
            copy(worksheet.cell(data_start_row, column).border)
            for column in range(1, template_column_count + 1)
        ]

        if worksheet.max_row >= data_start_row:
            worksheet.delete_rows(
                data_start_row,
                worksheet.max_row - data_start_row + 1,
            )

        def process_value(_index, value):
            if value and " - " in str(value):
                parts = str(value).split(" - ", 1)
                if len(parts) == 2 and parts[0].strip().isdigit():
                    value = parts[0].strip()
            return value

        for row_offset, submission in enumerate(submissions):
            data_dict = json.loads(submission.data_json)
            target_row = data_start_row + row_offset

            for column, cell_border in enumerate(template_borders, start=1):
                worksheet.cell(target_row, column).border = copy(cell_border)
            target_dimension = worksheet.row_dimensions[target_row]
            target_dimension.height = None
            target_dimension.hidden = False
            target_dimension.outlineLevel = 0
            target_dimension.collapsed = False

            for key, value in data_dict.items():
                if key.startswith('col_'):
                    column_index = int(key.split('_')[1])
                    worksheet.cell(
                        row=target_row,
                        column=column_index + 1,
                        value=process_value(column_index, value),
                    )

        workbook.save(download_path)
        return download_path
    except (KeyError, OSError, TypeError, ValueError) as exc:
        raise ExcelExportError(f"Không thể tạo file Excel: {exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()
