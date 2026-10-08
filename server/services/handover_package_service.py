"""Đóng gói bàn giao (lát G2): chép file theo kế hoạch chuẩn hóa (G1), kiểm SHA-256, xuất metadata NN-SIP.

- Chỉ đóng gói hồ sơ "sẵn sàng": mọi file trong thư mục hồ sơ không còn vấn đề ở kế hoạch G1 và mã
  hồ sơ không trùng với hồ sơ khác.
- Chép từ kho PDF (PDF_STORAGE_PATH) sang HANDOVER_DIR/<đường dẫn bàn giao>: ghi file tạm ".part"
  rồi đổi tên; SHA-256 phải khớp lúc tải lên. Chạy lại: file đích đã đúng SHA-256 thì bỏ qua.
- Kết quả trong HANDOVER_DIR/<gốc>/: Metadata_NN-SIP.xlsx (Metadata_HS, MetadataVB), SHA256SUMS.txt
  (kiểm bằng `sha256sum -c` trong thư mục gốc), Nhat_ky_dong_goi.xlsx.
- Trạng thái: HANDOVER_DIR/_jobs/project_<id>.json; khóa project_<id>.lock (mỗi dự án một lần đóng gói).
  Không có HANDOVER_DIR thì dùng thư mục "handover" cạnh kho PDF.
"""
import hashlib
import json
import logging
import os
import re
import threading
import unicodedata
from datetime import datetime, timezone
from pathlib import Path

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font

from server.repositories.normalization_repository import NormalizationRepository
from server.services.normalization_plan_service import build_plan

logger = logging.getLogger("server.handover")

HS_FIELDS = (
    "fileCode", "title", "maintenance", "mode", "language", "startDate", "endDate", "keyword",
    "totalDoc", "numberOfPaper", "numberOfPage", "format", "inforSign", "confidenceLevel",
    "paperFileCode", "riskRecovery", "riskRecoveryStatus", "description",
)
VB_FIELDS = (
    "fileCode", "docId", "docCode", "maintenance", "typeName", "codeNumber", "codeNotation", "issuedDate",
    "organName", "subject", "language", "numberOfPage", "inforSign", "keyword", "mode", "confidenceLevel",
    "autograph", "format", "riskRecovery", "riskRecoveryStatus", "process", "description",
)
SUPPLEMENT_FIELDS = ("Tệp tin hồ sơ", "Mục lục số", "Hộp số", "Hồ sơ số", "Tên phông", "Mã phông", "Giai đoạn/Nhiệm kỳ", "Path")
VB_EXTRA_FIELDS = ("Người ký", "File gốc")
# QC-01/QC-13: giá trị mặc định của dự án số hóa.
DEFAULTS = {"mode": "01", "language": "01", "confidenceLevel": "02", "format": "Bình thường"}
# Nhãn cột biểu mẫu (không dấu, chữ thường) -> trường metadata văn bản. So khớp nguyên nhãn.
VB_LABEL_ALIASES = {
    "typeName": ("the loai van ban", "ten loai van ban", "ten loai", "the loai", "loai van ban"),
    "codeNumber": ("so cua van ban", "so van ban", "so"),
    "codeNotation": ("ky hieu cua van ban", "ky hieu van ban", "ky hieu"),
    "issuedDate": ("ngay ky ngay thang nam van ban", "ngay ky", "ngay van ban", "ngay thang nam van ban", "ngay ban hanh"),
    "organName": ("ten co quan to chuc ban hanh van ban", "co quan ban hanh", "ten co quan ban hanh", "co quan to chuc ban hanh"),
    "subject": ("trich yeu noi dung", "trich yeu"),
    "keyword": ("tu khoa",),
    "autograph": ("but tich",),
    "description": ("ghi chu",),
    "Người ký": ("nguoi ky",),
}
UPPERCASE_FIELDS = {"typeName", "organName"}  # QC-01: IN HOA
ACTIVE_STATES = {"queued", "running"}
STATUS_EVERY_FILES = 20
_CHUNK = 1024 * 1024


class PackageFileError(Exception):
    pass


def handover_dir() -> Path:
    configured = os.getenv("HANDOVER_DIR")
    if configured:
        return Path(configured).resolve()
    return (Path(os.getenv("PDF_STORAGE_PATH", "uploads")).resolve().parent / "handover").resolve()


def storage_root() -> Path:
    return Path(os.getenv("PDF_STORAGE_PATH", "uploads")).resolve()


def job_status_path(project_id: int) -> Path:
    return handover_dir() / "_jobs" / f"project_{int(project_id)}.json"


def _lock_path(project_id: int) -> Path:
    return handover_dir() / "_jobs" / f"project_{int(project_id)}.lock"


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def read_job(project_id: int) -> dict | None:
    try:
        return json.loads(job_status_path(project_id).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None


def _write_job(project_id: int, payload: dict) -> dict:
    path = job_status_path(project_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {**payload, "updated_at": _now()}
    temporary = path.with_name(path.name + ".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8")
    os.replace(temporary, path)
    return payload


def _process_is_running(pid) -> bool:
    from server.services.export_job_service import _process_is_running as is_running

    return is_running(pid)


def _acquire_lock(project_id: int) -> None:
    path = _lock_path(project_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    for _attempt in range(2):
        try:
            descriptor = os.open(path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
        except FileExistsError:
            try:
                owner = int(path.read_text(encoding="utf-8").strip() or 0)
            except (OSError, ValueError):
                owner = 0
            if _process_is_running(owner):
                raise HTTPException(status_code=409, detail="Dự án đang được đóng gói, hãy đợi lần đóng gói này xong.")
            path.unlink(missing_ok=True)
            continue
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            handle.write(str(os.getpid()))
        return
    raise HTTPException(status_code=409, detail="Dự án đang được đóng gói, hãy đợi lần đóng gói này xong.")


def _release_lock(project_id: int) -> None:
    _lock_path(project_id).unlink(missing_ok=True)


def _normalize_label(text: str) -> str:
    value = str(text or "").replace("đ", "d").replace("Đ", "D")
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii").lower()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", value).split())


def vb_field_columns(schema) -> dict[str, str]:
    """Trường metadata văn bản -> tên ô (col_N) trong dữ liệu nhập, dò theo nhãn cột biểu mẫu."""
    alias_to_field = {alias: field for field, aliases in VB_LABEL_ALIASES.items() for alias in aliases}
    columns: dict[str, str] = {}
    for category in schema or []:
        for item in category.get("fields", []) if isinstance(category, dict) else []:
            field = alias_to_field.get(_normalize_label(item.get("label")))
            if field and field not in columns and item.get("name"):
                columns[field] = item["name"]
    return columns


def clean_text(value) -> str:
    """QC-15: bỏ khoảng trắng thừa, chuẩn Unicode NFC."""
    return " ".join(unicodedata.normalize("NFC", str(value or "")).split())


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(_CHUNK), b""):
            digest.update(chunk)
    return digest.hexdigest()


def copy_verified(source: Path, destination: Path, expected_sha256: str, expected_size: int) -> str:
    """Chép có kiểm SHA-256. Trả "skipped" nếu đích đã đúng, "copied" nếu vừa chép."""
    if not source.is_file():
        raise PackageFileError("Không tìm thấy file trong kho PDF của hệ thống")
    if destination.is_file() and destination.stat().st_size == expected_size and _sha256(destination) == expected_sha256:
        return "skipped"
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(destination.name + ".part")
    digest = hashlib.sha256()
    try:
        with source.open("rb") as reader, temporary.open("wb") as writer:
            for chunk in iter(lambda: reader.read(_CHUNK), b""):
                digest.update(chunk)
                writer.write(chunk)
        if digest.hexdigest() != expected_sha256:
            raise PackageFileError("SHA-256 của file trong kho khác lúc tải lên (file hỏng hoặc bị thay)")
        os.replace(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)
    return "copied"


def page_count(path: Path) -> int | None:
    try:
        import pypdf

        with path.open("rb") as handle:
            return len(pypdf.PdfReader(handle, strict=False).pages)
    except Exception:  # noqa: BLE001 - số trang chỉ để ghi metadata
        return None


def collect_package(db, project_id: int) -> dict:
    """Đọc mọi thứ cần cho đóng gói (một lần, rồi đóng kết nối CSDL)."""
    plan = build_plan(db, project_id=project_id, with_sources=True)
    repository = NormalizationRepository(db)
    project = repository.get(project_id)
    try:
        schema = json.loads(project.form_schema_json_snapshot or "[]")
    except (TypeError, ValueError):
        schema = []
    columns = vb_field_columns(schema)

    groups: dict[tuple, list] = {}
    unassigned = []
    for row in plan["rows"]:
        if row["dossier_code"] and row["dossier"] is not None:
            groups.setdefault((row["box"], row["folder"]), []).append(row)
        else:
            unassigned.append(row)
    codes: dict[str, list] = {}
    for key, rows in groups.items():
        codes.setdefault(rows[0]["dossier_code"], []).append(key)

    ready, skipped = [], []
    for key, rows in sorted(groups.items(), key=lambda item: (item[0][0] or 0, item[1][0]["dossier_code"])):
        reasons = sorted({problem for row in rows for problem in row["problems"]})
        if len(codes[rows[0]["dossier_code"]]) > 1:
            reasons.append("Trùng mã hồ sơ với thư mục khác")
        if reasons:
            skipped.append({"box": key[0], "folder": key[1], "dossier_code": rows[0]["dossier_code"], "reasons": reasons})
        else:
            ready.append(rows)
    unassigned_problems: dict[tuple, set] = {}
    for row in unassigned:
        unassigned_problems.setdefault((row["box"], row["folder"]), set()).update(row["problems"])
    for (box, folder), problems in sorted(unassigned_problems.items(), key=lambda item: (item[0][0] or 0, item[0][1])):
        skipped.append({"box": box, "folder": folder, "dossier_code": "", "reasons": sorted(problems)})

    document_ids = {row["asset"].assigned_document_id for rows in ready for row in rows if row["asset"].assigned_document_id}
    raw = repository.latest_submission_data_by_document(document_ids)
    entries = {}
    for document_id, data_json in raw.items():
        try:
            entries[document_id] = json.loads(data_json or "{}")
        except (TypeError, ValueError):
            entries[document_id] = {}
    return {
        "project": plan["project"],
        "root": plan["root"],
        "columns": columns,
        "ready": [
            [
                {
                    "relative_path": row["relative_path"],
                    "target_path": row["target_path"],
                    "kind": row["kind"],
                    "number": row["number"],
                    "dossier_code": row["dossier_code"],
                    "document_code": row["document_code"],
                    "storage_filename": row["asset"].storage_filename,
                    "sha256": row["asset"].content_sha256,
                    "size": row["asset"].byte_size,
                    "entry": entries.get(row["asset"].assigned_document_id, {}),
                    "box": row["box"],
                    "folder": row["folder"],
                    "dossier": {
                        "title": row["dossier"].title,
                        "maintenance": row["dossier"].maintenance_code,
                        "start_date": row["dossier"].start_date,
                        "end_date": row["dossier"].end_date,
                        "sheet_count": row["dossier"].sheet_count,
                        "catalog_number": row["dossier"].catalog_number,
                        "dossier_number": f"{row['dossier'].dossier_number}{row['dossier'].dossier_suffix or ''}",
                        "fonds_name": row["dossier"].fonds_name,
                        "fonds_code": row["dossier"].fonds_code,
                        "term": row["dossier"].term or "",
                        "note": row["dossier"].note or "",
                    },
                }
                for row in rows
            ]
            for rows in ready
        ],
        "skipped": skipped,
    }


def _vb_row(item: dict, columns: dict, pages: int | None) -> list:
    values = {field: "" for field in VB_FIELDS + SUPPLEMENT_FIELDS + VB_EXTRA_FIELDS}
    values.update(DEFAULTS)
    for field, column in columns.items():
        text = clean_text(item["entry"].get(column, ""))
        values[field] = text.upper() if field in UPPERCASE_FIELDS else text
    dossier = item["dossier"]
    values.update({
        "fileCode": item["dossier_code"], "docId": item["number"], "docCode": item["document_code"],
        "maintenance": dossier["maintenance"], "numberOfPage": pages if pages is not None else "",
        "Mục lục số": dossier["catalog_number"], "Hộp số": item["box"], "Hồ sơ số": dossier["dossier_number"],
        "Tên phông": dossier["fonds_name"], "Mã phông": dossier["fonds_code"], "Giai đoạn/Nhiệm kỳ": dossier["term"],
        "Path": item["target_path"], "File gốc": item["relative_path"],
    })
    return [values[field] for field in VB_FIELDS + SUPPLEMENT_FIELDS + VB_EXTRA_FIELDS]


def _hs_row(items: list, pages: dict) -> list:
    dossier = items[0]["dossier"]
    documents = [item for item in items if item["kind"] != "Bìa"]
    covers = [item for item in items if item["kind"] == "Bìa"]
    counted = [pages.get(item["target_path"]) for item in items]
    values = {field: "" for field in HS_FIELDS + SUPPLEMENT_FIELDS}
    values.update(DEFAULTS)
    values.update({
        "fileCode": items[0]["dossier_code"], "title": clean_text(dossier["title"]), "maintenance": dossier["maintenance"],
        "startDate": dossier["start_date"], "endDate": dossier["end_date"], "totalDoc": len(documents),
        "numberOfPaper": dossier["sheet_count"],
        "numberOfPage": sum(counted) if all(count is not None for count in counted) else "",
        "description": clean_text(dossier["note"]),
        "Tệp tin hồ sơ": covers[0]["target_path"].rsplit("/", 1)[-1] if covers else "",
        "Mục lục số": dossier["catalog_number"], "Hộp số": items[0]["box"], "Hồ sơ số": dossier["dossier_number"],
        "Tên phông": dossier["fonds_name"], "Mã phông": dossier["fonds_code"], "Giai đoạn/Nhiệm kỳ": dossier["term"],
        "Path": items[0]["target_path"].rsplit("/", 1)[0],
    })
    return [values[field] for field in HS_FIELDS + SUPPLEMENT_FIELDS]


def _sheet(workbook, title, header, rows, *, first=False):
    sheet = workbook.active if first else workbook.create_sheet(title)
    sheet.title = title
    sheet.append(list(header))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    return sheet


def _write_outputs(root_dir: Path, context: dict, packaged: list, pages: dict, results: list, failed: list) -> None:
    metadata = Workbook()
    _sheet(metadata, "Metadata_HS", HS_FIELDS + SUPPLEMENT_FIELDS, [_hs_row(items, pages) for items in packaged], first=True)
    _sheet(metadata, "MetadataVB", VB_FIELDS + SUPPLEMENT_FIELDS + VB_EXTRA_FIELDS, [
        _vb_row(item, context["columns"], pages.get(item["target_path"]))
        for items in packaged for item in items if item["kind"] != "Bìa"
    ])
    metadata.save(root_dir / "Metadata_NN-SIP.xlsx")

    prefix = f"{context['root']}/"
    lines = [f"{row['sha256']}  {row['target_path'][len(prefix):]}" for row in results if row["result"] != "Lỗi"]
    (root_dir / "SHA256SUMS.txt").write_text("\n".join(lines) + ("\n" if lines else ""), encoding="utf-8")

    log = Workbook()
    _sheet(log, "Đã đóng gói", ("File hiện tại", "Đường dẫn bàn giao", "SHA-256", "Dung lượng (byte)", "Số trang", "Kết quả"), [
        [row["relative_path"], row["target_path"], row["sha256"], row["size"], row["pages"], row["result"]]
        for row in results if row["result"] != "Lỗi"
    ], first=True)
    _sheet(log, "Chưa đóng gói", ("Hộp", "Thư mục hồ sơ", "Mã hồ sơ", "Lý do"), [
        [item["box"], item["folder"], item["dossier_code"], "; ".join(item["reasons"])]
        for item in context["skipped"] + failed
    ])
    _sheet(log, "Lỗi chép file", ("File hiện tại", "Đường dẫn bàn giao", "Lỗi"), [
        [row["relative_path"], row["target_path"], row["error"]] for row in results if row["result"] == "Lỗi"
    ])
    log.save(root_dir / "Nhat_ky_dong_goi.xlsx")


def run_package(project_id: int, session_factory, *, requested_by: int | None = None) -> dict:
    """Chạy đóng gói (trong luồng nền). Luôn nhả khóa và ghi trạng thái cuối."""
    job = read_job(project_id) or {}
    job.update({"project_id": project_id, "state": "running", "pid": os.getpid(), "message": "Đang đọc kế hoạch chuẩn hóa"})
    job = _write_job(project_id, job)
    try:
        with session_factory() as db:
            context = collect_package(db, project_id)
        base = handover_dir()
        root_dir = base / context["root"]
        root_dir.mkdir(parents=True, exist_ok=True)
        files = [item for items in context["ready"] for item in items]
        job.update({
            "root": context["root"], "output_dir": str(root_dir), "files_total": len(files), "files_done": 0,
            "files_copied": 0, "files_skipped": 0, "bytes_done": 0, "error_count": 0,
            "dossiers_ready": len(context["ready"]), "dossiers_skipped": len(context["skipped"]),
            "message": f"Đang chép {len(files)} file của {len(context['ready'])} hồ sơ",
        })
        job = _write_job(project_id, job)
        results, pages, errors = [], {}, []
        failed_codes = set()
        for index, item in enumerate(files, start=1):
            destination = base / item["target_path"]
            result = {**{key: item[key] for key in ("relative_path", "target_path", "sha256", "size")}, "pages": None}
            try:
                outcome = copy_verified(storage_root() / item["storage_filename"], destination, item["sha256"], item["size"])
                result["pages"] = page_count(destination)
                pages[item["target_path"]] = result["pages"]
                result["result"] = "Đã chép" if outcome == "copied" else "Đã có sẵn, đúng SHA-256"
                job["files_copied" if outcome == "copied" else "files_skipped"] += 1
                job["bytes_done"] += item["size"]
            except (PackageFileError, OSError) as error:
                result.update({"result": "Lỗi", "error": str(error)})
                failed_codes.add(item["dossier_code"])
                errors.append(f"{item['relative_path']}: {error}")
            results.append(result)
            job["files_done"] = index
            if index % STATUS_EVERY_FILES == 0:
                job = _write_job(project_id, job)
        packaged = [items for items in context["ready"] if items[0]["dossier_code"] not in failed_codes]
        failed = [
            {"box": items[0]["box"], "folder": items[0]["folder"], "dossier_code": items[0]["dossier_code"],
             "reasons": ["Lỗi chép file (xem sheet Lỗi chép file)"]}
            for items in context["ready"] if items[0]["dossier_code"] in failed_codes
        ]
        _write_outputs(root_dir, context, packaged, pages, results, failed)
        job.update({
            "state": "done", "finished_at": _now(), "dossiers_packaged": len(packaged),
            "error_count": len(errors), "errors": errors[:50],
            "message": (
                f"Đã đóng gói {len(packaged)} hồ sơ ({job['files_copied']} file mới chép, "
                f"{job['files_skipped']} file đã có sẵn) vào {root_dir}."
                + (f" {len(errors)} file lỗi." if errors else "")
                + (f" {len(context['skipped'])} hồ sơ chưa sẵn sàng." if context["skipped"] else "")
            ),
        })
        return _write_job(project_id, job)
    except Exception as error:  # noqa: BLE001 - mọi lỗi phải hiện ra trạng thái
        logger.exception("Đóng gói dự án %s thất bại", project_id)
        job.update({"state": "error", "finished_at": _now(), "message": f"Đóng gói thất bại: {error}"})
        return _write_job(project_id, job)
    finally:
        _release_lock(project_id)


def start_package(db, *, project_id: int, current_user: dict, session_factory, runner=None) -> dict:
    """Kiểm điều kiện rồi chạy đóng gói ở luồng nền. runner thay luồng nền trong test."""
    if NormalizationRepository(db).get(project_id) is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy dự án")
    _acquire_lock(project_id)
    try:
        context = collect_package(db, project_id)
        if not context["ready"]:
            raise HTTPException(
                status_code=409,
                detail="Chưa có hồ sơ nào sẵn sàng đóng gói. Tải Kế hoạch chuẩn hóa để xem cột Vấn đề.",
            )
        job = _write_job(project_id, {
            "project_id": project_id, "state": "queued", "pid": os.getpid(), "started_at": _now(),
            "requested_by": current_user.get("username"), "root": context["root"],
            "dossiers_ready": len(context["ready"]), "dossiers_skipped": len(context["skipped"]),
            "files_total": sum(len(items) for items in context["ready"]), "files_done": 0,
            "message": "Đang xếp hàng đóng gói",
        })
    except BaseException:
        _release_lock(project_id)
        raise
    if runner is None:
        thread = threading.Thread(
            target=run_package, args=(project_id, session_factory), name=f"handover-{project_id}", daemon=True
        )
        try:
            thread.start()
        except BaseException:
            _release_lock(project_id)
            raise
    else:
        runner(project_id, session_factory)
    return job


def fail_stale_package_jobs() -> int:
    """Khởi động máy chủ: lần đóng gói đang chạy mà tiến trình đã chết thì đánh lỗi, nhả khóa."""
    jobs_dir = handover_dir() / "_jobs"
    if not jobs_dir.is_dir():
        return 0
    failed = 0
    for path in jobs_dir.glob("project_*.json"):
        try:
            job = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
        if job.get("state") in ACTIVE_STATES and not _process_is_running(job.get("pid")):
            job.update({
                "state": "error", "finished_at": _now(),
                "message": "Máy chủ khởi động lại khi đang đóng gói. Bấm Đóng gói lại: file đã chép đúng sẽ được bỏ qua.",
            })
            project_id = int(path.stem.split("_")[-1])
            _write_job(project_id, job)
            _release_lock(project_id)
            failed += 1
    return failed
