"""Đối soát R1–R4 (lát R1, chỉ đọc). Giả định reviewer, chờ đối chiếu BA: BA gốc có R1–R4 nhưng không có trong repo.

- R1 Mục lục ↔ thư mục hồ sơ của file nhập liệu (theo hộp + số hồ sơ + hậu tố).
- R2 Gói S `done` mới nhất của hộp ↔ file nhập liệu của hộp (đường dẫn tính từ thư mục hộp, chỉ PDF,
  không phân biệt hoa thường, `\\` và `/` như nhau).
- R3 File nhập liệu (trừ bìa) ↔ hồ sơ nhập (submissions theo văn bản).
- R4 Số trang gói S `done` mới nhất ↔ tổng số tờ mục lục của hộp: hợp lý khi tờ ≤ trang ≤ 2 × tờ + số hồ sơ
  (giấy hai mặt, cộng thêm bìa mỗi hồ sơ).
"""
from io import BytesIO

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from server.repositories.reconciliation_repository import ReconciliationRepository
from server.services.arrangement_catalog_parser import (
    CatalogValueError,
    box_number_from_folder,
    box_number_of_case_key,
    parse_dossier_number,
)
from server.services.normalization_plan_service import STATUS_LABELS, is_cover_file, natural_key

OK = "Khớp"
PASSED = "Đạt"
REASONABLE = "Hợp lý"
GOOD_RESULTS = {OK, PASSED, REASONABLE}

R1_MISSING_FOLDER = "Thiếu thư mục"
R1_EXTRA_FOLDER = "Thừa thư mục"
R2_MISSING = "Thiếu ở nhập liệu"
R2_EXTRA = "Thừa ở nhập liệu"
R2_SIZE = "Khác dung lượng"
R2_NO_SCAN = "Hộp chưa có gói scan"
R3_NOT_ENTERED = "Chưa nhập"
R4_FEW_PAGES = "Ít trang hơn số tờ"
R4_MANY_PAGES = "Nhiều hơn 2 lần số tờ"
R4_NO_DATA = "Thiếu dữ liệu"

SHEETS = (
    ("R1", "R1 Mục lục-Thư mục", ("Hộp", "Hồ sơ số", "Tiêu đề (mục lục)", "Thư mục", "Số file", "Kết quả", "Ghi chú")),
    ("R2", "R2 Scan-Nhập liệu", ("Hộp", "Đường dẫn (từ thư mục hộp)", "Dung lượng scan", "Dung lượng nhập liệu", "Kết quả")),
    ("R3", "R3 File-Hồ sơ nhập", ("Hộp", "Đường dẫn file", "Số hồ sơ nhập", "Trạng thái", "Kết quả")),
    ("R4", "R4 Trang-Tờ", ("Hộp", "Số hồ sơ mục lục", "Tổng số tờ", "Số trang scan", "Ngưỡng trên (2 × tờ + hồ sơ)", "Kết quả")),
)
SUMMARY_TITLES = {
    "R1": "Mục lục ↔ thư mục hồ sơ",
    "R2": "Gói scan ↔ file nhập liệu",
    "R3": "File nhập liệu ↔ hồ sơ nhập",
    "R4": "Số trang scan ↔ số tờ mục lục",
}


def _parts(path) -> list[str]:
    return [part for part in str(path or "").replace("\\", "/").split("/") if part]


def _path_key(path) -> str:
    return "/".join(_parts(path)).casefold()


def _dossier_label(number: int, suffix: str) -> str:
    return f"{number}{suffix or ''}"


def _box_number(case_row) -> int | None:
    return box_number_of_case_key(case_row.case_key) or box_number_from_folder(case_row.display_name)


def _r1(cases, dossiers_by_case, assets_by_case, case_level):
    rows = []
    for case_row in cases:
        box = _box_number(case_row)
        catalog = {
            (dossier.dossier_number, (dossier.dossier_suffix or "").lower()): dossier
            for dossier in dossiers_by_case.get(case_row.id, [])
        }
        folders: dict[str, int] = {}
        for asset in assets_by_case.get(case_row.id, []):
            parts = _parts(asset.relative_path)
            if len(parts) == case_level + 2:
                folders[parts[case_level]] = folders.get(parts[case_level], 0) + 1
        matched = set()
        folder_rows = []
        for folder, count in folders.items():
            try:
                key = parse_dossier_number(folder)
            except CatalogValueError:
                folder_rows.append((_r1_order(None, folder), [box, "", "", folder, count, R1_EXTRA_FOLDER,
                                                              "Tên thư mục không đọc được số hồ sơ"]))
                continue
            dossier = catalog.get(key)
            if dossier is None:
                folder_rows.append((_r1_order(key, folder), [box, _dossier_label(*key), "", folder, count, R1_EXTRA_FOLDER, ""]))
            else:
                matched.add(key)
                folder_rows.append((_r1_order(key, folder), [box, _dossier_label(*key), dossier.title, folder, count, OK, ""]))
        for key, dossier in catalog.items():
            if key not in matched:
                folder_rows.append((_r1_order(key, ""), [box, _dossier_label(*key), dossier.title, "", 0, R1_MISSING_FOLDER, ""]))
        rows.extend(row for _order, row in sorted(folder_rows, key=lambda item: item[0]))
    return rows


def _r1_order(key, folder):
    """Theo số hồ sơ rồi hậu tố; thư mục không đọc được số xuống cuối."""
    if key is None:
        return (1, 0, "", natural_key(folder))
    return (0, key[0], key[1], natural_key(folder))


def _box_relative(asset, case_level) -> str:
    parts = _parts(asset.relative_path)
    return "/".join(parts[case_level:]) if len(parts) > case_level else "/".join(parts)


def _r2(cases, assets_by_case, packages, files_by_package, case_level):
    rows = []
    for case_row in cases:
        box = _box_number(case_row)
        assets = [asset for asset in assets_by_case.get(case_row.id, []) if asset.relative_path.lower().endswith(".pdf")]
        package = packages.get(case_row.id)
        if package is None:
            if assets:
                rows.append([box, "", None, None, R2_NO_SCAN])
            continue
        scanned = {}
        for scan_file in files_by_package.get(package.id, []):
            if scan_file.relative_path.lower().endswith(".pdf"):
                scanned[_path_key(scan_file.relative_path)] = scan_file
        entered = {_path_key(_box_relative(asset, case_level)): asset for asset in assets}
        case_rows = []
        for key in set(scanned) | set(entered):
            scan_file, asset = scanned.get(key), entered.get(key)
            shown = scan_file.relative_path.replace("\\", "/") if scan_file else _box_relative(asset, case_level)
            if asset is None:
                result = R2_MISSING
            elif scan_file is None:
                result = R2_EXTRA
            elif scan_file.file_size != asset.byte_size:
                result = R2_SIZE
            else:
                result = OK
            case_rows.append([box, shown, scan_file.file_size if scan_file else None,
                              asset.byte_size if asset else None, result])
        rows.extend(sorted(case_rows, key=lambda row: natural_key(row[1])))
    return rows


def _r3(cases, assets_by_case, submissions):
    rows = []
    for case_row in cases:
        box = _box_number(case_row)
        case_rows = []
        for asset in assets_by_case.get(case_row.id, []):
            filename = _parts(asset.relative_path)[-1] if _parts(asset.relative_path) else asset.original_filename
            if is_cover_file(filename):
                continue
            items = submissions.get(asset.assigned_document_id, []) if asset.assigned_document_id else []
            if not items:
                status, result = STATUS_LABELS[None], R3_NOT_ENTERED
            elif len(items) > 1:
                status = ", ".join(STATUS_LABELS.get(item[1], item[1]) for item in items)
                result = f"Nhiều hồ sơ ({len(items)})"
            else:
                status = STATUS_LABELS.get(items[0][1], items[0][1])
                result = PASSED if items[0][1] == "completed" else f"Chưa hoàn thành ({status})"
            case_rows.append([box, asset.relative_path, len(items), status, result])
        rows.extend(sorted(case_rows, key=lambda row: natural_key(row[1])))
    return rows


def _r4(cases, dossiers_by_case, packages, files_by_package):
    rows = []
    for case_row in cases:
        box = _box_number(case_row)
        dossiers = dossiers_by_case.get(case_row.id, [])
        sheets = sum(dossier.sheet_count or 0 for dossier in dossiers)
        package = packages.get(case_row.id)
        pages = None
        if package is not None:
            pages = sum(
                scan_file.page_count for scan_file in files_by_package.get(package.id, [])
                if scan_file.relative_path.lower().endswith(".pdf") and scan_file.page_count >= 0
            )
        upper = 2 * sheets + len(dossiers)
        if pages is None or not dossiers:
            result = R4_NO_DATA
        elif pages < sheets:
            result = R4_FEW_PAGES
        elif pages > upper:
            result = R4_MANY_PAGES
        else:
            result = REASONABLE
        rows.append([box, len(dossiers), sheets if dossiers else None, pages, upper if dossiers else None, result])
    return rows


def build_reconciliation(db, *, project_id: int) -> dict:
    repository = ReconciliationRepository(db)
    project = repository.get(project_id)
    if project is None:
        raise HTTPException(status_code=404, detail={"code": "project_not_found", "message": "Không tìm thấy dự án"})

    cases = sorted(
        repository.cases(project_id),
        key=lambda case_row: (_box_number(case_row) is None, _box_number(case_row) or 0, natural_key(case_row.display_name or "")),
    )
    dossiers_by_case: dict[int, list] = {}
    for dossier in repository.active_dossiers(project_id):
        dossiers_by_case.setdefault(dossier.case_id, []).append(dossier)
    assets_by_case: dict[int, list] = {}
    assets = repository.active_assets(project_id)
    for asset in assets:
        assets_by_case.setdefault(asset.case_id, []).append(asset)
    packages = repository.latest_done_packages(project_id)
    files_by_package: dict[int, list] = {}
    for scan_file in repository.scan_files([package.id for package in packages.values()]):
        files_by_package.setdefault(scan_file.package_id, []).append(scan_file)
    submissions = repository.submissions_by_document(
        {asset.assigned_document_id for asset in assets if asset.assigned_document_id is not None}
    )

    rows = {
        "R1": _r1(cases, dossiers_by_case, assets_by_case, project.case_level),
        "R2": _r2(cases, assets_by_case, packages, files_by_package, project.case_level),
        "R3": _r3(cases, assets_by_case, submissions),
        "R4": _r4(cases, dossiers_by_case, packages, files_by_package),
    }
    summary = {}
    for code, _title, columns in SHEETS:
        result_index = columns.index("Kết quả")
        counts: dict[str, int] = {}
        for row in rows[code]:
            counts[row[result_index]] = counts.get(row[result_index], 0) + 1
        summary[code] = {"title": SUMMARY_TITLES[code], "rows": len(rows[code]), "results": counts}
    return {"project": {"id": project.id, "name": project.name}, "summary": summary, "rows": rows}


def reconciliation_workbook(report: dict) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Tổng hợp"
    summary.append(["Dự án", report["project"]["name"]])
    summary.append(["Ghi chú", "Đối soát R1–R4: giả định reviewer, chờ đối chiếu BA"])
    summary.append([])
    summary.append(["Mã", "Đối soát", "Kết quả", "Số dòng"])
    for cell in summary[4]:
        cell.font = Font(bold=True)
    for code, info in report["summary"].items():
        if not info["results"]:
            summary.append([code, info["title"], "(không có dòng)", 0])
        for result, count in sorted(info["results"].items()):
            summary.append([code, info["title"], result, count])
    for column, width in zip("ABCD", (8, 34, 34, 10)):
        summary.column_dimensions[column].width = width

    header_fill = PatternFill("solid", fgColor="D9E2F3")
    warn = PatternFill("solid", fgColor="FCE4D6")
    for code, title, columns in SHEETS:
        sheet = workbook.create_sheet(title)
        sheet.append(list(columns))
        for index, name in enumerate(columns, start=1):
            cell = sheet.cell(1, index)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            sheet.column_dimensions[get_column_letter(index)].width = max(12, min(60, len(name) + 6))
        result_index = columns.index("Kết quả")
        for row in report["rows"][code]:
            sheet.append(row)
            if row[result_index] not in GOOD_RESULTS:
                for cell in sheet[sheet.max_row]:
                    cell.fill = warn
        sheet.freeze_panes = "A2"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
