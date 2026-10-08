"""Chi trả theo sản lượng (P1): đơn giá và bảng tạm tính (không lưu; P2 mới chốt kỳ).

Sản lượng theo người trong khoảng ngày (giờ Việt Nam, gồm cả hai đầu) — QC-07, QC-10; phần còn lại là giả định:
- NL (nhập liệu): văn bản nộp kiểm tra lần đầu trong kỳ (baseline chất lượng, như sheet Chấm công); người = người nhập.
- CN (check nhập liệu): văn bản duyệt lần đầu (`review_confirmed` đầu tiên) trong kỳ; người = người duyệt.
- SC-A4 (scan): trang A4 quy đổi của gói S `done` xong trong kỳ; người = người scan đã chọn khi Nộp S (trống:
  dòng "Chưa xác định người scan").
- Loại 2 (giấy xấu): văn bản thuộc hồ sơ mục lục có `bad_paper`; gói S thuộc hộp có ít nhất 1 hồ sơ `bad_paper`.
  Đơn giá loại 2 = đơn giá loại 1 × `bad_paper_factor` của Chính sách dự án (QC-07, mặc định 1,3).
- Thiếu đơn giá: thành tiền trống kèm cảnh báo, không lỗi.
"""
from datetime import date, datetime, time, timedelta, timezone
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from io import BytesIO

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font

from server.database import get_utc_now
from server.models_payroll import ProjectWorkRate
from server.repositories.payroll_repository import PayrollRepository
from server.repositories.timesheet_repository import TimesheetRepository
from server.services.account_policy_service import VIETNAM_TZ
from server.services.arrangement_catalog_parser import CatalogValueError, parse_dossier_number
from server.services.normalization_plan_service import is_cover_file
from server.services.project_policy_service import QC_VERSION, get_effective_policy
from server.services.timesheet_service import vietnam_date

MAX_PERIOD_DAYS = 92
RATE_CODES = ("NL-1", "CN-1", "SC-A4-1")
WORKS = {
    "NL": ("Nhập liệu", "văn bản"),
    "CN": ("Check nhập liệu", "văn bản"),
    "SC-A4": ("Scan A4", "trang A4"),
}
UNKNOWN_SCANNER = "Chưa xác định người scan"


def _error(status, code, message):
    return HTTPException(status_code=status, detail={"code": code, "message": message})


def _plain(value: Decimal | None):
    if value is None:
        return None
    return int(value) if value == value.to_integral_value() else float(value)


def _project_or_404(repository, project_id):
    project = repository.get(project_id)
    if project is None:
        raise _error(404, "project_not_found", "Không tìm thấy dự án")
    return project


def _rate_rows(repository, project_id):
    rows = repository.rates(project_id)
    return {code: rows[code].unit_price for code in RATE_CODES if code in rows}


def get_rates(db, *, project_id: int) -> dict:
    repository = PayrollRepository(db)
    _project_or_404(repository, project_id)
    rows = repository.rates(project_id)
    policy = get_effective_policy(db, project_id=project_id)
    return {
        "rates": [
            {
                "code": code,
                "label": WORKS[code.rsplit("-", 1)[0]][0],
                "unit": WORKS[code.rsplit("-", 1)[0]][1],
                "unit_price": _plain(rows[code].unit_price) if code in rows else None,
            }
            for code in RATE_CODES
        ],
        "bad_paper_factor": _plain(Decimal(str(policy["bad_paper_factor"]))),
    }


def update_rates(db, *, project_id: int, rates: dict, actor: dict) -> dict:
    """rates: {mã: số ≥ 0 hoặc null (xóa đơn giá)}; chỉ nhận NL-1, CN-1, SC-A4-1."""
    repository = PayrollRepository(db)
    _project_or_404(repository, project_id)
    cleaned = {}
    for code, value in (rates or {}).items():
        if code not in RATE_CODES:
            raise _error(400, "unknown_work_code", f"Mã công việc không hợp lệ: {code}")
        if value is None or value == "":
            cleaned[code] = None
            continue
        if isinstance(value, bool):
            raise _error(400, "invalid_rate", f"Đơn giá {code} phải là số")
        try:
            price = Decimal(str(value))
        except InvalidOperation:
            raise _error(400, "invalid_rate", f"Đơn giá {code} phải là số")
        if not price.is_finite() or price < 0:
            raise _error(400, "negative_rate", f"Đơn giá {code} không được âm")
        if price >= Decimal("1e12"):
            raise _error(400, "invalid_rate", f"Đơn giá {code} quá lớn")
        cleaned[code] = price.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    existing = repository.rates(project_id)
    now = get_utc_now()
    for code, price in cleaned.items():
        row = existing.get(code)
        if price is None:
            if row is not None:
                repository.delete_rate(row)
        elif row is None:
            repository.add_rate(ProjectWorkRate(project_id=project_id, work_code=code, unit_price=price,
                                                updated_by_user_id=actor["id"], updated_at=now))
        else:
            row.unit_price = price
            row.updated_by_user_id = actor["id"]
            row.updated_at = now
    db.commit()
    return get_rates(db, project_id=project_id)


def parse_period(date_from: str, date_to: str) -> tuple[date, date]:
    try:
        start = date.fromisoformat(str(date_from))
        end = date.fromisoformat(str(date_to))
    except ValueError:
        raise _error(400, "invalid_date", "Ngày phải có dạng YYYY-MM-DD")
    if start > end:
        raise _error(400, "period_invalid", "Từ ngày phải trước hoặc bằng đến ngày")
    if (end - start).days + 1 > MAX_PERIOD_DAYS:
        raise _error(400, "period_too_long", f"Mỗi kỳ tối đa {MAX_PERIOD_DAYS} ngày")
    return start, end


def _utc_start(day: date) -> datetime:
    return datetime.combine(day, time.min, tzinfo=VIETNAM_TZ).astimezone(timezone.utc).replace(tzinfo=None)


def _bad_paper_lookup(repository, project):
    """(văn bản giấy xấu?, hộp giấy xấu?) theo mục lục: hồ sơ = thư mục ngay dưới thư mục hộp."""
    dossiers = {}
    bad_cases = set()
    for dossier in repository.active_dossiers(project.id):
        dossiers[(dossier.case_id, dossier.dossier_number, (dossier.dossier_suffix or "").lower())] = dossier
        if dossier.bad_paper:
            bad_cases.add(dossier.case_id)

    def asset_is_bad(asset) -> bool:
        parts = [part for part in str(asset.relative_path or "").replace("\\", "/").split("/") if part]
        if len(parts) != project.case_level + 2:
            return False
        try:
            number, suffix = parse_dossier_number(parts[project.case_level])
        except CatalogValueError:
            return False
        dossier = dossiers.get((asset.case_id, number, suffix))
        return bool(dossier and dossier.bad_paper)

    return asset_is_bad, bad_cases


def compute_payroll(db, *, project_id: int, date_from: date, date_to: date) -> dict:
    repository = PayrollRepository(db)
    project = _project_or_404(repository, project_id)
    policy = get_effective_policy(db, project_id=project_id)
    factor = Decimal(str(policy["bad_paper_factor"]))
    rates = _rate_rows(repository, project_id)
    asset_is_bad, bad_cases = _bad_paper_lookup(repository, project)

    quantities: dict[tuple, int] = {}

    def add(person_id, code, amount):
        if amount:
            key = (person_id, code)
            quantities[key] = quantities.get(key, 0) + amount

    in_period = lambda day: day is not None and date_from <= day <= date_to  # noqa: E731
    documents = {}
    for asset in repository.active_assets(project_id):
        filename = str(asset.relative_path or asset.original_filename).replace("\\", "/").rsplit("/", 1)[-1]
        if not is_cover_file(filename):
            documents[asset.assigned_document_id] = asset_is_bad(asset)
    submissions = [s for s in repository.latest_submissions_by_document(documents).values() if s.status != "draft"]
    timesheet = TimesheetRepository(db)
    ids = [submission.id for submission in submissions]
    assessments = timesheet.assessments_by_submission(ids)
    reviews = timesheet.first_review_confirmations(ids)
    for submission in submissions:
        kind = "2" if documents.get(submission.assigned_document_id) else "1"
        input_user_id, submitted_at = assessments.get(submission.id, (None, None))
        if in_period(vietnam_date(submitted_at or submission.created_at)):
            add(input_user_id or submission.created_by_user_id, f"NL-{kind}", 1)
        review = reviews.get(submission.id)
        if review is not None and in_period(vietnam_date(review[1])):
            add(review[0], f"CN-{kind}", 1)

    for package in repository.done_scan_packages_between(project_id, _utc_start(date_from), _utc_start(date_to + timedelta(days=1))):
        kind = "2" if package.case_id in bad_cases else "1"
        add(package.scanned_by_user_id, f"SC-A4-{kind}", package.total_a4_equivalent or 0)

    names = timesheet.display_names({person for person, _code in quantities if person is not None})

    def name_of(person_id, code):
        if person_id is None:
            return UNKNOWN_SCANNER if code.startswith("SC-") else "Chưa xác định người"
        return names.get(person_id) or f"Tài khoản #{person_id}"

    lines = []
    missing = set()
    for (person_id, code), quantity in quantities.items():
        work, kind = code.rsplit("-", 1)
        base_code = f"{work}-1"
        price = rates.get(base_code)
        line_factor = factor if kind == "2" else Decimal("1")
        amount = None
        if price is None:
            missing.add(base_code)
        else:
            amount = (Decimal(quantity) * price * line_factor).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
        lines.append({
            "user_id": person_id, "name": name_of(person_id, code), "work_code": code, "work": WORKS[work][0],
            "unit": WORKS[work][1], "quantity": quantity, "unit_price": _plain(price), "factor": _plain(line_factor),
            "amount": _plain(amount),
        })
    lines.sort(key=lambda line: (line["user_id"] is None, line["name"].casefold(), line["work_code"]))

    people = {}
    for line in lines:
        person = people.setdefault(line["user_id"], {"user_id": line["user_id"], "name": line["name"], "amount": 0, "complete": True})
        if line["amount"] is None:
            person["complete"] = False
        else:
            person["amount"] += line["amount"]
    return {
        "project": {"id": project.id, "name": project.name},
        "period": {"from": date_from.isoformat(), "to": date_to.isoformat()},
        "params": {
            "rates": {code: _plain(price) for code, price in rates.items()},
            "bad_paper_factor": _plain(factor),
            "qc_version": QC_VERSION,
        },
        "lines": lines,
        "people": list(people.values()),
        "total": sum(line["amount"] or 0 for line in lines),
        "warnings": [f"Chưa có đơn giá {code}" for code in RATE_CODES if code in missing],
    }


def payroll_workbook(result: dict, *, title: str = "Bảng tạm tính chi trả") -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Tổng theo người"
    period = result["period"]
    params = result["params"]
    header = [
        [title, result["project"]["name"]],
        ["Kỳ", f"{period['from']} đến {period['to']} (giờ Việt Nam)"],
        ["Đơn giá loại 1", ", ".join(f"{code}: {price}" for code, price in params["rates"].items()) or "(chưa có)"],
        ["Hệ số giấy xấu", params["bad_paper_factor"]],
        ["Quy chuẩn", params["qc_version"]],
    ] + [["Cảnh báo", warning] for warning in result["warnings"]]
    for row in header:
        summary.append(row)
    summary.append([])
    summary.append(["Người", "Thành tiền (VNĐ)", "Ghi chú"])
    for cell in summary[summary.max_row]:
        cell.font = Font(bold=True)
    for person in result["people"]:
        summary.append([person["name"], person["amount"], "" if person["complete"] else "Thiếu đơn giá một số dòng"])
    summary.append(["Tổng", result["total"], ""])
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 40

    detail = workbook.create_sheet("Chi tiết")
    detail.append(["Người", "Mã công việc", "Công việc", "Đơn vị", "Sản lượng", "Đơn giá", "Hệ số", "Thành tiền"])
    for cell in detail[1]:
        cell.font = Font(bold=True)
    for line in result["lines"]:
        detail.append([line["name"], line["work_code"], line["work"], line["unit"], line["quantity"],
                       line["unit_price"], line["factor"], line["amount"]])
    for column, width in zip("ABCDEFGH", (32, 12, 18, 10, 12, 14, 8, 16)):
        detail.column_dimensions[column].width = width
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
