"""Read and check an arrangement catalogue (mục lục chỉnh lý) workbook.

Pure functions, no database access (FR-ARR-01; QC-05, QC-07, QC-13, QC-16).
One row of sheet ``Muc_luc`` is one dossier (hồ sơ); columns are matched by
their header text, so their order does not matter.
"""

import calendar
import hashlib
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime

MAX_CATALOG_BYTES = 5 * 1024 * 1024
MAX_CATALOG_ROWS = 10_000
MAX_REPORTED_ERRORS = 500
CATALOG_SHEET = "Muc_luc"
FIRST_VALID_YEAR = 1945

# QC-13: thời hạn bảo quản.
MAINTENANCE_CODES = {"01": "Vĩnh viễn", "02": "Có thời hạn (lâu dài)"}

# Boxes created from the catalogue before their scan folder exists. The ":"
# can never come from an uploaded folder path, so the key cannot collide.
PLACEHOLDER_BOX_KEY_PREFIX = "::muc-luc/hop-"


@dataclass(frozen=True)
class CatalogColumn:
    key: str
    label: str
    required: bool


CATALOG_COLUMNS = (
    CatalogColumn("stt", "STT", False),
    CatalogColumn("fonds_code", "Mã phông", True),
    CatalogColumn("fonds_name", "Tên phông", True),
    CatalogColumn("catalog_number", "Mục lục số", True),
    CatalogColumn("box_number", "Hộp số", True),
    CatalogColumn("dossier", "Hồ sơ số", True),
    CatalogColumn("file_notation", "Ký hiệu hồ sơ", False),
    CatalogColumn("title", "Tiêu đề hồ sơ", True),
    CatalogColumn("start_date", "Thời gian bắt đầu", True),
    CatalogColumn("end_date", "Thời gian kết thúc", True),
    CatalogColumn("maintenance_code", "Thời hạn bảo quản", True),
    CatalogColumn("sheet_count", "Số tờ", True),
    CatalogColumn("term", "Giai đoạn/Nhiệm kỳ", False),
    CatalogColumn("bad_paper", "Giấy xấu", False),
    CatalogColumn("note", "Ghi chú", False),
)
_COLUMN_LABELS = {column.key: column.label for column in CATALOG_COLUMNS}
_TEXT_LIMITS = {
    "fonds_code": 50,
    "fonds_name": 255,
    "catalog_number": 50,
    "title": 1000,
    "term": 100,
    "note": 1000,
}
# Dossier comparison fields (everything but the key and the source row).
DOSSIER_FIELDS = (
    "fonds_code", "fonds_name", "catalog_number", "file_notation", "title",
    "start_date", "end_date", "start_year", "maintenance_code", "sheet_count",
    "term", "bad_paper", "note",
)

_DATE_TEXT = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})$")
_DOSSIER = re.compile(r"^0*(\d{1,6})([A-Za-z]?)$")
_POSITIVE_INTEGER = re.compile(r"^0*(\d{1,9})$")
_FILE_NOTATION = re.compile(r"^[A-Za-z0-9_-]{1,20}$")


class CatalogValueError(ValueError):
    pass


@dataclass
class CatalogRow:
    row: int
    box_number: int
    dossier_number: int
    dossier_suffix: str
    values: dict

    @property
    def key(self):
        return (self.box_number, self.dossier_number, self.dossier_suffix)

    @property
    def dossier_label(self):
        return f"{self.dossier_number}{self.dossier_suffix}"


@dataclass
class ParsedCatalog:
    file_sha256: str
    rows: list = field(default_factory=list)
    errors: list = field(default_factory=list)  # {"row", "column", "message"}
    file_errors: list = field(default_factory=list)
    data_rows: int = 0

    @property
    def error_count(self):
        return len(self.errors) + len(self.file_errors)


def normalize_header(value) -> str:
    text = unicodedata.normalize("NFC", str(value or ""))
    text = text.replace("(*)", "")
    return re.sub(r"\s+", " ", text).strip().casefold()


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return unicodedata.normalize("NFC", str(value)).strip()


def parse_positive_integer(value, label) -> int:
    if isinstance(value, bool):
        raise CatalogValueError(f"{label} phải là số nguyên dương.")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    match = _POSITIVE_INTEGER.match(_text(value))
    if not match or int(match.group(1)) < 1:
        raise CatalogValueError(f"{label} phải là số nguyên dương (ví dụ 20 hoặc 0020).")
    return int(match.group(1))


def parse_dossier_number(value) -> tuple[int, str]:
    """``12``, ``0012``, ``12a``, ``12B`` -> (12, "") / (12, "a") / (12, "b")."""
    if isinstance(value, bool):
        raise CatalogValueError("Hồ sơ số phải là số nguyên, có thể kèm 1 chữ cái (ví dụ 12, 12a).")
    match = _DOSSIER.match(_text(value))
    if not match or int(match.group(1)) < 1:
        raise CatalogValueError("Hồ sơ số phải là số nguyên, có thể kèm 1 chữ cái (ví dụ 12, 12a).")
    return int(match.group(1)), match.group(2).lower()


def parse_catalog_date(value, *, today: date) -> tuple[int, int, int]:
    """QC-05 date as ``(year, month, day)``; 0 marks an unknown day or month."""
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        day, month, year = value.day, value.month, value.year
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        raise CatalogValueError("Ngày phải ghi dạng dd/mm/yyyy hoặc là ô ngày của Excel.")
    else:
        match = _DATE_TEXT.match(_text(value))
        if not match:
            raise CatalogValueError(
                "Ngày phải ghi dạng dd/mm/yyyy (thiếu ngày: 00/mm/yyyy, thiếu ngày tháng: 00/00/yyyy)."
            )
        day, month, year = (int(part) for part in match.groups())
    if not FIRST_VALID_YEAR <= year <= today.year:
        raise CatalogValueError(f"Năm phải từ {FIRST_VALID_YEAR} đến {today.year}.")
    if month > 12:
        raise CatalogValueError("Tháng phải từ 01 đến 12 (00 nếu không rõ).")
    if day > 31:
        raise CatalogValueError("Ngày phải từ 01 đến 31 (00 nếu không rõ).")
    if month == 0 and day != 0:
        raise CatalogValueError("Không rõ tháng thì ngày cũng ghi 00 (00/00/yyyy).")
    if day and month:
        try:
            date(year, month, day)
        except ValueError:
            raise CatalogValueError("Ngày không tồn tại trong lịch.") from None
    return year, month, day


def format_catalog_date(parts) -> str:
    year, month, day = parts
    return f"{day:02d}/{month:02d}/{year:04d}"


def _earliest(parts):
    year, month, day = parts
    return year, month or 1, day or 1


def _latest(parts):
    year, month, day = parts
    month = month or 12
    return year, month, day or calendar.monthrange(year, month)[1]


def _maintenance_code(value) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = _text(value)
    if text.isdigit() and len(text) == 1:
        text = text.zfill(2)
    if text not in MAINTENANCE_CODES:
        codes = ", ".join(MAINTENANCE_CODES)
        raise CatalogValueError(f"Thời hạn bảo quản phải là mã trong danh mục QC-13 ({codes}).")
    return text


def _bad_paper(value) -> bool:
    text = _text(value)
    if not text:
        return False
    if text.casefold() == "x":
        return True
    raise CatalogValueError("Giấy xấu chỉ ghi x hoặc để trống.")


def _limited_text(key, value, *, required) -> str | None:
    text = _text(value)
    label = _COLUMN_LABELS[key]
    if not text:
        if required:
            raise CatalogValueError(f"{label} là cột bắt buộc.")
        return None
    limit = _TEXT_LIMITS.get(key)
    if limit and len(text) > limit:
        raise CatalogValueError(f"{label} dài quá {limit} ký tự.")
    return text


def _parse_row(cells, *, row_number, today):
    """Return ``(CatalogRow | None, [errors])`` for one worksheet row."""
    errors = []
    values = {}

    def check(key, parser):
        try:
            return parser()
        except CatalogValueError as error:
            errors.append({"row": row_number, "column": _COLUMN_LABELS[key], "message": str(error)})
            return None

    def required_cell(key):
        value = cells.get(key)
        if _text(value) == "":
            raise CatalogValueError(f"{_COLUMN_LABELS[key]} là cột bắt buộc.")
        return value

    box = check("box_number", lambda: parse_positive_integer(required_cell("box_number"), "Hộp số"))
    dossier = check("dossier", lambda: parse_dossier_number(required_cell("dossier")))
    for key in ("fonds_code", "fonds_name", "catalog_number", "title"):
        values[key] = check(key, lambda key=key: _limited_text(key, cells.get(key), required=True))
    for key in ("term", "note"):
        values[key] = check(key, lambda key=key: _limited_text(key, cells.get(key), required=False))

    def notation():
        text = _text(cells.get("file_notation"))
        if text and not _FILE_NOTATION.match(text):
            raise CatalogValueError(
                "Ký hiệu hồ sơ chỉ gồm chữ không dấu, số và _ - (tối đa 20 ký tự)."
            )
        return text or None

    values["file_notation"] = check("file_notation", notation)
    start = check("start_date", lambda: parse_catalog_date(required_cell("start_date"), today=today))
    end = check("end_date", lambda: parse_catalog_date(required_cell("end_date"), today=today))
    if start and end and _earliest(start) > _latest(end):
        errors.append({
            "row": row_number,
            "column": _COLUMN_LABELS["start_date"],
            "message": "Thời gian bắt đầu sau thời gian kết thúc.",
        })
    values["maintenance_code"] = check(
        "maintenance_code", lambda: _maintenance_code(required_cell("maintenance_code"))
    )
    values["sheet_count"] = check(
        "sheet_count", lambda: parse_positive_integer(required_cell("sheet_count"), "Số tờ")
    )
    values["bad_paper"] = check("bad_paper", lambda: _bad_paper(cells.get("bad_paper")))
    if errors:
        return None, errors
    values["start_date"] = format_catalog_date(start)
    values["end_date"] = format_catalog_date(end)
    values["start_year"] = start[0]
    return CatalogRow(row_number, box, dossier[0], dossier[1], values), []


def _row_is_blank(cells):
    return all(_text(value) == "" for value in cells.values())


def parse_catalog_workbook(content: bytes, *, today: date) -> ParsedCatalog:
    parsed = ParsedCatalog(file_sha256=hashlib.sha256(content).hexdigest())
    if not content.startswith(b"PK"):
        parsed.file_errors.append("Chỉ nhận file Excel .xlsx (theo mẫu QC-16).")
        return parsed
    if len(content) > MAX_CATALOG_BYTES:
        parsed.file_errors.append("File mục lục vượt quá 5 MB.")
        return parsed
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError, TypeError):
        parsed.file_errors.append("Không đọc được file. Hãy lưu lại dạng .xlsx theo mẫu QC-16.")
        return parsed

    try:
        sheet = next(
            (ws for ws in workbook.worksheets if ws.title.casefold() == CATALOG_SHEET.casefold()),
            None,
        )
        if sheet is None:
            parsed.file_errors.append(f"Không tìm thấy sheet {CATALOG_SHEET}.")
            return parsed
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None) or ()
        by_header = {normalize_header(value): index for index, value in enumerate(header) if value is not None}
        positions = {}
        missing = []
        for column in CATALOG_COLUMNS:
            index = by_header.get(normalize_header(column.label))
            if index is None:
                if column.required:
                    missing.append(column.label)
                continue
            positions[column.key] = index
        if missing:
            parsed.file_errors.append("Thiếu cột bắt buộc: " + ", ".join(missing) + ".")
            return parsed

        seen = {}
        data_rows = 0
        for row_number, raw in enumerate(rows, start=2):
            cells = {
                key: (raw[index] if index < len(raw) else None)
                for key, index in positions.items()
            }
            if _row_is_blank(cells):
                continue
            data_rows += 1
            parsed.data_rows = data_rows
            if data_rows > MAX_CATALOG_ROWS:
                parsed.file_errors.append(
                    f"File có hơn {MAX_CATALOG_ROWS:,} hồ sơ; hãy tách thành nhiều file.".replace(",", ".")
                )
                parsed.rows.clear()
                return parsed
            catalog_row, errors = _parse_row(cells, row_number=row_number, today=today)
            parsed.errors.extend(errors)
            if catalog_row is None:
                continue
            first_row = seen.get(catalog_row.key)
            if first_row is not None:
                parsed.errors.append({
                    "row": row_number,
                    "column": _COLUMN_LABELS["dossier"],
                    "message": (
                        f"Trùng Hộp số {catalog_row.box_number} + Hồ sơ số "
                        f"{catalog_row.dossier_label} với dòng {first_row}."
                    ),
                })
                continue
            seen[catalog_row.key] = row_number
            parsed.rows.append(catalog_row)
        if data_rows == 0:
            parsed.file_errors.append("Sheet Muc_luc chưa có hồ sơ nào (dữ liệu bắt đầu từ dòng 2).")
    finally:
        workbook.close()
    return parsed


# --- Folder structure: <Hộp>\<Hồ sơ>\*.pdf from the project's case level ---


def box_number_from_folder(name) -> int | None:
    match = _POSITIVE_INTEGER.match(str(name or "").strip())
    if not match or int(match.group(1)) < 1:
        return None
    return int(match.group(1))


def placeholder_box_key(box_number: int) -> str:
    return f"{PLACEHOLDER_BOX_KEY_PREFIX}{int(box_number)}"


def is_placeholder_box_key(case_key) -> bool:
    return str(case_key or "").startswith(PLACEHOLDER_BOX_KEY_PREFIX)


def box_number_of_case_key(case_key) -> int | None:
    key = str(case_key or "")
    if is_placeholder_box_key(key):
        return box_number_from_folder(key[len(PLACEHOLDER_BOX_KEY_PREFIX):])
    return box_number_from_folder(key.rstrip("/").rsplit("/", 1)[-1])


def catalog_structure_problems(relative_paths, *, case_level, existing_box_keys=None, examples=3):
    """Why these PDFs do not follow <Hộp>\\<Hồ sơ>\\*.pdf from ``case_level``.

    The folder at ``case_level`` is the box (named by its number) and the next
    one is the dossier, which holds the PDFs. ``existing_box_keys`` maps a box
    number to the case key already used for it in the project.
    """
    too_shallow, too_deep, not_numbered = [], [], []
    box_keys = {}
    for path in relative_paths:
        text = str(path).replace("\\", "/").strip("/")
        folders = [part for part in text.split("/")[:-1] if part]
        if len(folders) == case_level:
            too_shallow.append(text)
            continue
        if len(folders) > case_level + 1:
            too_deep.append(text)
            continue
        if len(folders) < case_level:
            continue  # rejected earlier by the manifest checks
        box_folder = folders[case_level - 1]
        number = box_number_from_folder(box_folder)
        if number is None:
            not_numbered.append(text)
            continue
        box_keys.setdefault(number, set()).add("/".join(folders[:case_level]).casefold())

    def sample(paths):
        shown = "; ".join(paths[:examples])
        more = f" … (tổng {len(paths)} file)" if len(paths) > examples else ""
        return shown + more

    problems = []
    if too_shallow:
        problems.append(
            f"PDF nằm ngay trong thư mục cấp {case_level}, nên cấp hồ sơ của dự án đang là "
            f"thư mục HỒ SƠ chứ không phải HỘP: {sample(too_shallow)}. Dự án dùng mục lục cần "
            f"cấp hồ sơ = cấp thư mục Hộp (<Hộp>\\<Hồ sơ>\\*.pdf)."
        )
    if too_deep:
        problems.append(
            f"PDF nằm sâu hơn <Hộp>\\<Hồ sơ> tính từ cấp {case_level} (thư mục cấp {case_level} "
            f"có vẻ là PHÔNG, hoặc hồ sơ có thư mục con): {sample(too_deep)}."
        )
    if not_numbered:
        problems.append(
            f"Thư mục cấp {case_level} phải đặt tên bằng Hộp số (ví dụ 20 hoặc 0020): "
            f"{sample(not_numbered)}."
        )
    existing_box_keys = existing_box_keys or {}
    for number in sorted(box_keys):
        keys = set(box_keys[number])
        existing = existing_box_keys.get(number)
        if existing and not is_placeholder_box_key(existing):
            keys.add(existing.casefold())
        if len(keys) > 1:
            problems.append(f"Hộp {number} xuất hiện ở nhiều thư mục: {', '.join(sorted(keys))}.")
    return problems
