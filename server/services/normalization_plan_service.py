"""Kế hoạch chuẩn hóa (lát G1): mã hồ sơ, mã văn bản, đường dẫn bàn giao theo QC-03/QC-04.

Chỉ đọc: không đổi tên, không chép file. Admin tải Excel về để soát trước khi đóng gói (G2).
Nguồn: file nhập liệu đang dùng của dự án (`project_document_assets`, cấu trúc <Hộp>/<Hồ sơ>/*.pdf
tính từ cấp hộp), mục lục chỉnh lý (`arrangement_dossiers`) và chính sách dự án (mã cơ quan, ký hiệu).
"""
import re
import unicodedata
from io import BytesIO

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from server.repositories.normalization_repository import NormalizationRepository
from server.services.arrangement_catalog_parser import (
    CatalogValueError,
    box_number_from_folder,
    box_number_of_case_key,
    parse_dossier_number,
)
from server.services.project_policy_service import get_effective_policy

# QC-13: mã thời hạn bảo quản -> viết tắt trong cây thư mục bàn giao (QC-04).
MAINTENANCE_ABBREVIATIONS = {"01": "VV", "02": "LD"}
STATUS_LABELS = {
    None: "Chưa nhập",
    "draft": "Lưu nháp",
    "pending_review": "Chờ duyệt",
    "pending_input_confirmation": "Chờ người nhập xác nhận",
    "completed": "Hoàn thành",
}
PROBLEM_NO_ORGAN = "Chưa có Mã cơ quan (Chính sách dự án)"
PROBLEM_NOT_IN_DOSSIER = "File không nằm trong thư mục hồ sơ (<Hộp>/<Hồ sơ>/file.pdf)"
PROBLEM_BAD_FOLDER = "Tên thư mục hồ sơ không đọc được số hồ sơ"
PROBLEM_NOT_IN_CATALOG = "Thư mục hồ sơ không có trong mục lục"
PROBLEM_MAINTENANCE = "Thời hạn bảo quản không có trong danh mục QC-13"
PROBLEM_NOT_APPROVED = "Văn bản chưa hoàn thành nhập liệu"
PLAN_COLUMNS = (
    ("Hộp", 10), ("Thư mục hồ sơ", 14), ("File hiện tại", 40), ("Loại", 9), ("STT", 6),
    ("Mã hồ sơ", 26), ("Mã văn bản", 34), ("Đường dẫn bàn giao", 90), ("Nhập liệu", 22), ("Vấn đề", 60),
)
_UNSAFE = re.compile(r"[^A-Za-z0-9._-]+")
_NUMBER_PARTS = re.compile(r"(\d+)")


def ascii_name(text) -> str:
    """Tên thư mục/file theo QC-04: chữ không dấu, số, '.', '_', '-'; khoảng trắng thành '_'."""
    value = str(text or "").replace("đ", "d").replace("Đ", "D")
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    value = re.sub(r"\s+", "_", value.strip())
    return _UNSAFE.sub("", value)


def natural_key(name: str) -> list:
    """QC-12: thứ tự tự nhiên (2.pdf trước 10.pdf)."""
    return [int(part) if part.isdigit() else part.casefold() for part in _NUMBER_PARTS.split(name)]


def is_cover_file(filename: str) -> bool:
    """QC-03: file bìa (BIA.pdf hoặc tên chứa BIA) đi theo hồ sơ, không đánh STT."""
    return "bia" in ascii_name(filename).casefold()


def dossier_code(organ_code: str, start_year: int, number: int, suffix: str, notation: str | None) -> str:
    """QC-03: {Mã cơ quan}.{Năm}.{Số HS 2 chữ số}[.{Ký hiệu}]; hậu tố chữ (12a) giữ sau số."""
    code = f"{organ_code}.{start_year}.{number:02d}{suffix or ''}"
    return f"{code}.{notation}" if notation else code


def handover_root(project) -> str:
    return f"CSDL_SOHOA_{ascii_name(project.name) or project.id}"


def _case_box_number(case_row) -> int | None:
    return box_number_of_case_key(case_row.case_key) or box_number_from_folder(case_row.display_name)


def _location(asset, case_level: int) -> tuple[str | None, str]:
    """Thư mục hồ sơ (ngay dưới thư mục hộp) và tên file của một file nhập liệu."""
    parts = [part for part in str(asset.relative_path or "").replace("\\", "/").split("/") if part]
    filename = parts[-1] if parts else asset.original_filename
    if len(parts) == case_level + 2:
        return parts[case_level], filename
    return None, filename


def build_plan(db, *, project_id: int) -> dict:
    repository = NormalizationRepository(db)
    project = repository.get(project_id)
    if project is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy dự án")
    policy = get_effective_policy(db, project_id=project_id)
    organ_code = policy.get("organ_code") or ""
    default_notation = policy.get("file_notation") or None
    root = handover_root(project)

    cases = {case_row.id: case_row for case_row in repository.cases(project_id)}
    by_case: dict[int, dict] = {}
    by_box: dict[int, dict] = {}
    for dossier in repository.active_dossiers(project_id):
        key = (dossier.dossier_number, (dossier.dossier_suffix or "").lower())
        if dossier.case_id is not None:
            by_case.setdefault(dossier.case_id, {})[key] = dossier
        by_box.setdefault(dossier.box_number, {})[key] = dossier

    assets = repository.active_assets(project_id)
    statuses = repository.latest_submission_status_by_document(
        {asset.assigned_document_id for asset in assets if asset.assigned_document_id is not None}
    )
    groups: dict[tuple, list] = {}
    for asset in assets:
        folder, filename = _location(asset, project.case_level)
        groups.setdefault((asset.case_id, folder), []).append((filename, asset))

    def group_order(item):
        (case_id, folder), _files = item
        case_row = cases.get(case_id)
        box = _case_box_number(case_row) if case_row else None
        return (box is None, box or 0, natural_key(case_row.display_name if case_row else ""), natural_key(folder or ""))

    rows = []
    for (case_id, folder), files in sorted(groups.items(), key=group_order):
        case_row = cases.get(case_id)
        box = _case_box_number(case_row) if case_row else None
        common = []
        dossier = None
        if not organ_code:
            common.append(PROBLEM_NO_ORGAN)
        if folder is None:
            common.append(PROBLEM_NOT_IN_DOSSIER)
        else:
            try:
                key = parse_dossier_number(folder)
            except CatalogValueError:
                common.append(PROBLEM_BAD_FOLDER)
            else:
                dossier = by_case.get(case_id, {}).get(key) or by_box.get(box, {}).get(key)
                if dossier is None:
                    common.append(PROBLEM_NOT_IN_CATALOG)
        abbreviation = MAINTENANCE_ABBREVIATIONS.get(dossier.maintenance_code) if dossier else None
        if dossier is not None and abbreviation is None:
            common.append(PROBLEM_MAINTENANCE)
        code = ""
        base = ""
        if dossier is not None and organ_code:
            notation = ascii_name(dossier.file_notation) or default_notation
            code = dossier_code(organ_code, dossier.start_year, dossier.dossier_number,
                                (dossier.dossier_suffix or "").lower(), notation)
            if abbreviation:
                base = "/".join([root, organ_code, str(dossier.start_year), abbreviation,
                                 ascii_name(dossier.fonds_code), code])
        number = 0
        for filename, asset in sorted(files, key=lambda item: natural_key(item[0])):
            problems = list(common)
            status = statuses.get(asset.assigned_document_id)
            cover = is_cover_file(filename)
            if cover:
                document_code = ""
                target = f"{base}/{code}_BIA.pdf" if base else ""
            else:
                number += 1
                document_code = f"{code}.{number:07d}" if code else ""
                target = f"{base}/{document_code}.pdf" if base else ""
                if status != "completed":
                    problems.append(PROBLEM_NOT_APPROVED)
            rows.append({
                "box": box,
                "folder": folder or "",
                "relative_path": asset.relative_path,
                "kind": "Bìa" if cover else "Văn bản",
                "number": None if cover else number,
                "dossier_code": code,
                "document_code": document_code,
                "target_path": target,
                "entry_status": STATUS_LABELS.get(status, status),
                "problems": problems,
            })

    problem_counts: dict[str, int] = {}
    for row in rows:
        for problem in row["problems"]:
            problem_counts[problem] = problem_counts.get(problem, 0) + 1
    return {
        "project": {"id": project.id, "name": project.name},
        "organ_code": organ_code,
        "root": root,
        "summary": {
            "files": len(rows),
            "documents": sum(1 for row in rows if row["kind"] == "Văn bản"),
            "covers": sum(1 for row in rows if row["kind"] == "Bìa"),
            "dossiers": len({(row["box"], row["folder"]) for row in rows if row["dossier_code"]}),
            "ready": sum(1 for row in rows if not row["problems"]),
            "with_problems": sum(1 for row in rows if row["problems"]),
            "problems": problem_counts,
        },
        "rows": rows,
    }


def plan_workbook(plan: dict) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Tổng hợp"
    info = plan["summary"]
    lines = [
        ("Dự án", plan["project"]["name"]),
        ("Mã cơ quan", plan["organ_code"] or "(chưa có)"),
        ("Thư mục gốc bàn giao", plan["root"]),
        ("Số thư mục hồ sơ có mã", info["dossiers"]),
        ("Số văn bản", info["documents"]),
        ("Số file bìa", info["covers"]),
        ("Số file sẵn sàng", info["ready"]),
        ("Số file còn vấn đề", info["with_problems"]),
    ] + [(f"Vấn đề: {name}", count) for name, count in sorted(info["problems"].items())]
    for line in lines:
        summary.append(list(line))
    summary.column_dimensions["A"].width = 60
    summary.column_dimensions["B"].width = 40
    for cell in summary["A"]:
        cell.font = Font(bold=True)

    sheet = workbook.create_sheet("Kế hoạch đổi tên")
    sheet.append([name for name, _width in PLAN_COLUMNS])
    for index, (_name, width) in enumerate(PLAN_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        header = sheet.cell(1, index)
        header.font = Font(bold=True)
        header.fill = PatternFill("solid", fgColor="D9E2F3")
    warn = PatternFill("solid", fgColor="FCE4D6")
    for row in plan["rows"]:
        sheet.append([
            row["box"], row["folder"], row["relative_path"], row["kind"], row["number"],
            row["dossier_code"], row["document_code"], row["target_path"], row["entry_status"],
            "; ".join(row["problems"]),
        ])
        if row["problems"]:
            sheet.cell(sheet.max_row, len(PLAN_COLUMNS)).fill = warn
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
