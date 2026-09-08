"""Validate a built export worker using a NEW disposable PostgreSQL cluster.

Never accepts an existing database URL or data directory. Keeps test artifacts
under scratch and stops only the PostgreSQL cluster created by this invocation.
"""

from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
import secrets
import socket
import subprocess
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--postgres-bin", required=True, type=Path)
    parser.add_argument("--app", required=True, type=Path)
    args = parser.parse_args()
    app = args.app.resolve(strict=True)
    postgres_bin = args.postgres_bin.resolve(strict=True)
    scratch = ROOT / "scratch"
    scratch.mkdir(exist_ok=True)
    work = Path(tempfile.mkdtemp(prefix="export-release-smoke-", dir=scratch))
    data = work / "pgdata"
    with socket.socket() as listener:
        listener.bind(("127.0.0.1", 0))
        port = listener.getsockname()[1]
    flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)

    def pg(tool, *arguments):
        # A daemon may inherit stdout/stderr. Files avoid waiting for its pipe
        # handles to close after pg_ctl itself has already exited on Windows.
        log = work / f"{tool}.log"
        with log.open("ab") as output:
            result = subprocess.run(
                [str(postgres_bin / f"{tool}.exe"), *map(str, arguments)],
                stdin=subprocess.DEVNULL, stdout=output, stderr=output,
                timeout=60, creationflags=flags,
            )
        if result.returncode:
            raise RuntimeError(f"{tool}: {log.read_text(errors='replace')}")
        return result

    print(f"Isolated artifacts: {work}", flush=True)
    pg("initdb", "-D", data, "-U", "export_smoke", "-A", "trust", "--no-locale", "-E", "UTF8")
    engine = None
    started = False
    try:
        started = True
        pg("pg_ctl", "-D", data, "-l", work / "postgres.log", "-o",
           f"-h 127.0.0.1 -p {port}", "-w", "start")
        database_url = f"postgresql+psycopg://export_smoke@127.0.0.1:{port}/postgres"
        from server.host_runtime_paths import HostRuntimePaths

        paths = HostRuntimePaths(work / "host-data")
        paths.ensure_directories()
        config = {
            "APP_ENV": "production", "DATABASE_URL": database_url,
            "SECRET_KEY": secrets.token_urlsafe(48), "REDIS_URL": "",
            "PDF_STORAGE_PATH": str(paths.uploads_dir),
            "TEMPLATE_STORAGE_PATH": str(paths.templates_dir),
            "DOCUMENT_SOURCE_ROOT": str(paths.source_documents_dir),
            "EXPORT_WORK_DIR": str(paths.exports_dir), "LOG_DIR": str(paths.logs_dir),
        }
        paths.config_path.write_text(
            "\n".join(f"{key}='{value}'" for key, value in config.items()) + "\n",
            encoding="utf-8",
        )
        os.environ.update(config)
        os.environ["SCAN_TO_EXCEL_DATA_DIR"] = str(paths.root)
        from sqlalchemy import create_engine, text
        from sqlalchemy.orm import Session
        from openpyxl import Workbook, load_workbook
        from server.migration_runner import migration_state, upgrade_database
        from server.models import Submission, SubmissionReviewAssignment, Template, User
        from server.routers import submissions

        engine = create_engine(database_url)
        resources = app.parent / "_internal"
        upgrade_database(engine, base_dir=resources)
        upgrade_database(engine, base_dir=resources)
        assert migration_state(engine).ready
        print("PostgreSQL migration and repeat upgrade: PASS", flush=True)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        for row in (["Group", None], ["Name", "Value"], ["Label", "Label"], ["Detail", "Detail"]):
            sheet.append(row)
        workbook.save(paths.templates_dir / "smoke.xlsx")
        workbook.close()
        with Session(engine) as db:
            template = Template(name="Smoke", filename="smoke.xlsx")
            author = User(username="smoke_author", password="unused", full_name="Author")
            reviewer = User(username="smoke_reviewer", password="unused", full_name="Reviewer")
            db.add_all([template, author, reviewer])
            db.flush()
            row = Submission(
                template_id=template.id, created_by_user_id=author.id,
                data_json='{"col_0":"Smoke approved","col_1":"42"}',
                status="pending_review",
            )
            db.add(row)
            db.flush()
            db.add(SubmissionReviewAssignment(submission_id=row.id, reviewer_user_id=reviewer.id))
            db.commit()
            identity = {"id": reviewer.id, "role": "user", "username": reviewer.username}
            lease = submissions.api_claim_submission_view(row.id, current_user=identity, db=db)
            result = submissions.api_confirm_submission_review(
                row.id, submissions.ReviewContentRequest(data=json.loads(row.data_json)),
                lease_token=lease["lease_token"], current_user=identity, db=db,
            )
            assert result["submission_status"] == "completed"
            template_id = template.id
        print("PostgreSQL review confirmation: PASS", flush=True)

        job_id = secrets.token_hex(16)
        status_path = paths.exports_dir / f"export_job_{job_id}.json"
        status_path.write_text(json.dumps({"job_id": job_id, "state": "queued"}), encoding="utf-8")
        (paths.exports_dir / "export_all.lock").write_text(job_id, encoding="utf-8")
        result = subprocess.run(
            [str(app), "--export-worker", "--job-id", job_id,
             "--template-id", str(template_id), "--extension", ".xlsx"],
            cwd=work, env=os.environ.copy(), capture_output=True, text=True,
            errors="replace", timeout=60, creationflags=flags,
        )
        if result.returncode:
            raise RuntimeError(f"Frozen worker: {result.stdout}\n{result.stderr}")
        status = json.loads(status_path.read_text(encoding="utf-8"))
        assert status["state"] == "completed", status
        workbook = load_workbook(paths.exports_dir / f"export_job_{job_id}.xlsx")
        assert workbook["Data"]["A5"].value == "Smoke approved"
        assert workbook["Data"]["B5"].value == "42"
        workbook.close()
        assert not (paths.exports_dir / "export_all.lock").exists()
        print("Frozen worker and exported cell values: PASS", flush=True)

        connection = ["-h", "127.0.0.1", "-p", str(port), "-U", "export_smoke"]
        dump = work / "smoke.dump"
        pg("pg_dump", *connection, "-d", "postgres", "-Fc", "-f", dump)
        pg("createdb", *connection, "restored")
        pg("pg_restore", *connection, "-d", "restored", "--exit-on-error", dump)
        restored = create_engine(database_url.rsplit("/", 1)[0] + "/restored")
        try:
            assert migration_state(restored).ready
            with restored.connect() as db:
                row = db.execute(text("SELECT status, data_json FROM submissions")).one()
                assert row.status == "completed"
                assert json.loads(row.data_json)["col_0"] == "Smoke approved"
        finally:
            restored.dispose()
        print("PostgreSQL backup/restore and data verification: PASS", flush=True)
        return 0
    finally:
        if engine is not None:
            engine.dispose()
        if started and (data / "postmaster.pid").exists():
            pg("pg_ctl", "-D", data, "-m", "fast", "-w", "stop")
            print("Isolated PostgreSQL stopped; artifacts retained.", flush=True)


if __name__ == "__main__":
    raise SystemExit(main())
